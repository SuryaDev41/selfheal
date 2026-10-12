"""Focused tests for conversion evidence and generic browser assertions."""

from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook
from playwright.sync_api import sync_playwright

from selfheal.executor.browser import parse_test_data
from selfheal.executor.generic import GenericStepExecutor
from selfheal.healing.locator import HealingMatch, verified_locator
from selfheal.llm.openai_client import LocatorResponseError, _structured_response
from selfheal.onboarding import _convert_cases, _visible_discovery_requests
from selfheal.suite_converter import (
    _action_issue,
    _claim_issues,
    _scope_form_fields,
    converted_suite_path,
    model_cases,
    replace_reviewed_conversions,
    reviewed_attempts,
    reviewed_case_ids,
    save_converted_suite,
    validate_conversion,
)
from selfheal.suite_loader import read_excel_tests


def _case(case_id: str, expected: str, data: str = "") -> dict:
    return {"id": case_id, "title": case_id, "step": "Open home", "data": data,
            "expected": expected, "row": 2}


def _proposal(case_id: str, steps: list[str], assertions: list[str]) -> dict:
    return {"cases": [{"id": case_id, "steps": steps,
                       "checks": [{"clause_index": 0, "assertions": assertions}],
                       "review_reason": ""}]}


def test_discovery_defers_controls_missing_from_the_initial_page():
    snapshots = [{"page": "/login", "html": (
        '<h2>New User Signup!</h2><input placeholder="Name">'
        '<button>Signup</button><h2>Login to your account</h2>'
        '<input placeholder="Email Address"><button>Login</button>'
    )}]
    requests = [
        {"page": "/login", "command": "FILL", "target": "Name in New User Signup!"},
        {"page": "/login", "command": "CLICK", "target": "Login"},
        {"page": "/login", "command": "CLICK", "target": "Logout"},
        {"page": "/checkout", "command": "CLICK", "target": "Pay and Confirm Order"},
    ]
    assert _visible_discovery_requests(requests, snapshots) == requests[:2]


def test_add_to_cart_is_not_satisfied_by_opening_the_cart():
    source = "Login\nAdd product to cart and go to Cart\nClick 'Proceed To Checkout'"
    steps = ["CLICK | Login", "CLICK | Cart", "CLICK | Proceed To Checkout"]
    assert _action_issue(source, steps) == "Source add-to-cart action is missing from generated steps"


def test_generated_click_cannot_use_a_url_path_as_element_locator():
    source = [_case("TC_001", "'Done' title is displayed")]
    result = validate_conversion(
        source,
        _proposal("TC_001", ["OPEN | /", "CLICK | /done", "ASSERT_TITLE | Done"],
                  ["ASSERT_TITLE | Done"]),
        "https://example.test",
    )[0]
    assert "CLICK needs an element locator" in result["conversion_note"]


def test_mixed_suite_keeps_explicit_case_out_of_openai_conversion():
    explicit = {**_case("TC_001", ""), "step": "OPEN | /\nASSERT_TITLE | Home"}
    prose = _case("TC_002", "'Done' title is displayed")

    class Client:
        def convert_site(self, _pages, cases):
            assert [case["id"] for case in cases] == ["TC_002"]
            return _proposal("TC_002", ["OPEN | /done", "ASSERT_TITLE | Done"],
                             ["ASSERT_TITLE | Done"])

    converted, calls = _convert_cases(Client(), [], [explicit, prose], "https://example.test")
    assert calls == 1
    assert converted[0]["step"] == explicit["step"]
    assert converted[0]["conversion_note"] == ""
    assert converted[1]["conversion_note"] == ""


def test_whole_workbook_review_uses_one_call_and_preserves_explicit_steps():
    explicit = {**_case("TC_001", ""), "step": "OPEN | /\nASSERT_TITLE | Home"}
    prose = [_case(f"TC_{number:03}", "'Done' title is displayed") for number in range(2, 21)]

    class Client:
        calls = 0

        def convert_site(self, _pages, cases):
            self.calls += 1
            assert len(cases) == 20
            return {"cases": [
                {"id": case["id"], "steps": ["OPEN | /wrong"], "checks": [],
                 "review_reason": ""} if case["id"] == "TC_001" else
                {"id": case["id"], "steps": ["OPEN | /done", "ASSERT_TITLE | Done"],
                 "checks": [{"clause_index": 0, "assertions": ["ASSERT_TITLE | Done"]}],
                 "review_reason": ""}
                for case in cases
            ]}

    client = Client()
    converted, calls = _convert_cases(
        client, [], [explicit, *prose], "https://example.test",
        one_call=True, review_all=True,
    )

    assert calls == client.calls == 1
    assert len(converted) == 20
    assert converted[0]["step"] == explicit["step"]
    assert all(not case["conversion_note"] for case in converted)


def test_upload_review_does_not_rewrite_commands_or_source_workbook(tmp_path):
    source = tmp_path / "upload.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append(["TestCase ID", "Title", "Test Step", "Test Data", "Expected Result"])
    explicit = "OPEN | /\nASSERT_TEXT | css=h1 | 'Done'"
    sheet.append(["TC_001", "Existing commands", explicit, "", ""])
    sheet.append(["TC_002", "Prose", "Open the Done page", "", "'Done' title is displayed"])
    with_expected = "OPEN | /\nASSERT_TEXT | css=h1 | Done"
    sheet.append(["TC_003", "Existing commands with expected result", with_expected, "",
                  "'Done' heading is visible"])
    book.save(source)
    book.close()
    original_bytes = source.read_bytes()

    class ReviewModel:
        calls = 0

        def convert_site(self, pages, cases):
            self.calls += 1
            assert pages == [{"page": "/", "html": "<h1>Done</h1>"}]
            assert cases[0]["steps"] == explicit
            return {"cases": [
                {"id": "TC_001", "steps": ["OPEN | /wrong"], "checks": [], "review_reason": ""},
                {"id": "TC_002", "steps": ["OPEN | /done", "ASSERT_TITLE | Done"],
                 "checks": [{"clause_index": 0, "assertions": ["ASSERT_TITLE | Done"]}],
                 "review_reason": ""},
                {"id": "TC_003", "steps": ["OPEN | /wrong"],
                 "checks": [{"clause_index": 0, "assertions": ["ASSERT_TEXT | css=h1 | Done"]}],
                 "review_reason": ""},
            ]}

    model = ReviewModel()
    reviewed, calls = _convert_cases(
        model, [{"page": "/", "html": "<h1>Done</h1>"}], read_excel_tests(source),
        "https://example.test", one_call=True, review_all=True,
    )
    output = converted_suite_path(tmp_path, "https://example.test", source)
    output.parent.mkdir(parents=True)
    save_converted_suite(source, output, reviewed)

    assert calls == model.calls == 1
    assert source.read_bytes() == original_bytes
    assert output != source
    copied = read_excel_tests(output)
    assert copied[0]["step"] == explicit
    assert copied[1]["step"] == "OPEN | /done\nASSERT_TITLE | Done"
    assert copied[2]["step"] == with_expected
    assert all(not case["conversion_note"] for case in copied)


def test_website_review_note_is_kept_without_rewriting_explicit_steps():
    explicit = {**_case("TC_001", ""), "step": "OPEN | /\nASSERT_URL | /"}

    class ReviewModel:
        def convert_site(self, _pages, _cases):
            return {"cases": [{"id": "TC_001", "steps": ["OPEN | /wrong"],
                               "checks": [], "review_reason": "Site contradicts the expected page"}]}

    reviewed, calls = _convert_cases(
        ReviewModel(), [], [explicit], "https://example.test", one_call=True, review_all=True,
    )

    assert calls == 1
    assert reviewed[0]["step"] == explicit["step"]
    assert reviewed[0]["conversion_note"] == "Site contradicts the expected page"


def test_reviewed_retry_sends_prior_feedback_but_redacts_prior_values(tmp_path):
    path = tmp_path / "converted.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Framework Cases"
    sheet.append(["TestCase ID", "Test Step", "Conversion Note"])
    sheet.append(["TC_001", "FILL | email | qa@example.test\nASSERT_TEXT | message | Done",
                  "Expected Result clause has no executable assertion"])
    book.save(path)
    book.close()

    attempts = reviewed_attempts(path)
    case = {**_case("TC_001", "'Done' is displayed"), **attempts["TC_001"]}
    sent = model_cases([case])[0]
    assert sent["validation_feedback"] == attempts["TC_001"]["validation_feedback"]
    assert sent["previous_steps"] == ["FILL | email | <redacted>",
                                      "ASSERT_TEXT | message | Done"]


def test_conversion_reports_multiple_independent_defects():
    source = [_case("TC_001", "'Done' title is displayed; no products listed", "Email: <email>")]
    payload = {"cases": [{"id": "TC_001", "steps": [
        "OPEN | /", "FILL | email | ${DATA:email}", "ASSERT_TITLE | Done",
    ], "checks": [{"clause_index": 0, "assertions": ["ASSERT_TITLE | Done"]}],
        "review_reason": ""}]}
    result = validate_conversion(source, payload, "https://example.test")[0]
    note = result["conversion_note"]
    assert "Every Expected Result clause" in note
    assert "Clause 2: Expected Result clause" in note
    assert "real Test Data value: email" in note
    assert len(result["conversion_issues"]) >= 3


def test_timestamp_email_is_validated_after_runtime_expansion():
    steps = [
        "OPEN | /login",
        "FILL | css=input[type=email] | ${DATA:email}",
        "ASSERT_VALUE | css=input[type=email] | ${DATA:email}",
    ]
    source = [{**_case("TC_004", "", "Email: qa-<timestamp>@example.com"),
               "step": "\n".join(steps)}]
    payload = {"cases": [{"id": "TC_004", "steps": steps, "checks": [], "review_reason": ""}]}

    assert validate_conversion(source, payload, "https://example.test")[0]["conversion_note"] == ""

    source[0]["data"] = "Email: qa-<unknown>@example.com"
    assert "real Test Data value: email" in validate_conversion(
        source, payload, "https://example.test",
    )[0]["conversion_note"]
    source[0]["data"] = "Email: qa-<timestamp>"
    assert "real Test Data value: email" in validate_conversion(
        source, payload, "https://example.test",
    )[0]["conversion_note"]


def test_explicit_blank_login_validation_is_not_mistaken_for_form_entry():
    steps = [
        "OPEN | /login",
        "CLICK | css=button[data-qa='login-button']",
        "ASSERT_VALIDATION | css=input[data-qa='login-email'] | Please fill out this field",
        "ASSERT_URL | /login",
    ]
    source = [{**_case("TC_006", "", "Email: (blank)\nPassword: (blank)"),
               "step": "\n".join(steps)}]
    payload = {"cases": [{"id": "TC_006", "steps": steps, "checks": [], "review_reason": ""}]}

    assert validate_conversion(source, payload, "https://example.test")[0]["conversion_note"] == ""
    assert _action_issue("Enter email, then submit", ["CLICK | Submit"]) == (
        "Source form entry is missing from generated steps"
    )


def test_in_page_alert_class_does_not_require_browser_dialog_commands():
    steps = [
        "OPEN | /",
        "FILL | css=#subscribe | ${DATA:valid}",
        "CLICK | css=#subscribe-button",
        "ASSERT_TEXT | css=footer .alert-success.alert | Subscribed!",
    ]
    source = [{**_case("TC_019", "", "Valid: qa-<timestamp>@example.com"),
               "step": "\n".join(steps)}]
    payload = {"cases": [{"id": "TC_019", "steps": steps, "checks": [], "review_reason": ""}]}

    assert validate_conversion(
        source, payload, "https://example.test", allow_site_actions=True,
    )[0]["conversion_note"] == ""


def test_native_field_validation_does_not_require_an_in_page_message():
    assertion = "ASSERT_VALIDATION | css=input[type=email] | any"
    assert _claim_issues(
        "HTML5 required-field validation appears on Email",
        [assertion],
        ["OPEN | /login", assertion],
    ) == []


def test_server_login_error_cannot_be_verified_with_native_validation_only():
    assertion = "ASSERT_VALIDATION | css=input[type=email] | Your email or password is incorrect!"
    issues = _claim_issues(
        "Error 'Your email or password is incorrect!' is displayed",
        [assertion],
        ["OPEN | /login", assertion],
    )
    assert "Server-rendered error needs an in-page text or visibility assertion" in issues


def test_wait_visible_waits_for_a_delayed_direct_locator(tmp_path):
    executor = GenericStepExecutor(
        {"environments": {"prod": {"base_url": "https://example.test"}}},
        "prod", tmp_path / "suite.xlsx", 2000,
    )
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=True, channel="chromium")
        try:
            page = browser.new_page()
            page.set_content("<div id='result' hidden>Added!</div>")
            page.evaluate("setTimeout(() => document.querySelector('#result').hidden = false, 200)")
            assert executor.execute(page, "WAIT_VISIBLE | css=#result", {}) == (
                "Waited for css=#result to become visible"
            )
        finally:
            browser.close()


def test_focused_retry_keeps_partial_progress_without_marking_case_ready():
    source = [_case("TC_001", "'Done' title is displayed; no products listed")]

    class Client:
        calls = 0

        def convert_site(self, _pages, cases):
            self.calls += 1
            if self.calls == 1:
                return {"cases": [{"id": "TC_001", "steps": ["OPEN | /"],
                                   "checks": [], "review_reason": ""}]}
            assert "Every Expected Result clause" in cases[0]["validation_feedback"]
            return {"cases": [{"id": "TC_001", "steps": ["OPEN | /", "ASSERT_TITLE | Done"],
                               "checks": [
                                   {"clause_index": 0, "assertions": ["ASSERT_TITLE | Done"]},
                                   {"clause_index": 1, "assertions": []},
                               ], "review_reason": ""}]}

    client = Client()
    converted, calls = _convert_cases(client, [], source, "https://example.test")
    assert calls == client.calls == 2
    assert "ASSERT_TITLE | Done" in converted[0]["step"]
    assert converted[0]["conversion_note"]
    assert not converted[0]["execute_partial"]


def test_name_fields_are_derived_only_from_a_two_part_name():
    assert parse_test_data("Name: QA Tester")["first_name"] == "QA"
    assert parse_test_data("Name: QA Tester")["last_name"] == "Tester"
    assert "last_name" not in parse_test_data("Name: QA")
    assert parse_test_data("Name: QA Tester\nFirst_name: Custom")["first_name"] == "Custom"


def test_structured_response_rejects_incomplete_or_refused_output():
    incomplete = SimpleNamespace(choices=[SimpleNamespace(
        finish_reason="length", message=SimpleNamespace(content='{"cases": []}', refusal=None),
    )])
    with pytest.raises(LocatorResponseError, match="incomplete"):
        _structured_response(incomplete, "conversion")
    refused = SimpleNamespace(choices=[SimpleNamespace(
        finish_reason="stop", message=SimpleNamespace(content=None, refusal="declined"),
    )])
    with pytest.raises(LocatorResponseError, match="declined"):
        _structured_response(refused, "discovery")


def test_form_heading_from_source_scopes_repeated_fields_and_submit():
    source = "Enter name and email under 'New User Signup!'\nClick 'Signup'"
    steps = ["FILL | name | ${DATA:name}", "FILL | email | ${DATA:email}",
             "CLICK | button[type='submit']"]
    assert _scope_form_fields(source, steps) == [
        "FILL | name in New User Signup! | ${DATA:name}",
        "FILL | email in New User Signup! | ${DATA:email}",
        "CLICK | Signup in New User Signup!",
    ]


def test_multiple_assertions_cover_a_compound_expected_clause():
    source = [_case("TC_001", "'SEARCHED PRODUCTS' heading is shown with no products listed")]
    steps = ["OPEN | /products", "ASSERT_TEXT | css=h2 | SEARCHED PRODUCTS",
             "ASSERT_COUNT | css=.product-card | 0"]
    result = validate_conversion(
        source, _proposal("TC_001", steps, steps[1:]), "https://example.test",
    )
    assert not result[0]["conversion_note"]

    weak = validate_conversion(
        source, _proposal("TC_001", steps[:-1], steps[1:2]), "https://example.test",
    )
    assert "zero" in weak[0]["conversion_note"]


def test_generated_ui_quotes_are_not_treated_as_literal_page_text():
    source = [_case("TC_001", "'Done' heading is displayed")]
    step = "ASSERT_TEXT | css=h1 | 'Done'"
    converted = validate_conversion(
        source, _proposal("TC_001", ["OPEN | /done", step], [step]),
        "https://example.test",
    )
    assert converted[0]["conversion_note"] == ""
    assert converted[0]["step"].endswith("ASSERT_TEXT | css=h1 | Done")


def test_arithmetic_claim_requires_both_calculation_and_stated_values():
    source = [_case("TC_014", "Cart shows quantity 4 and Total = Rs. 2000 (4 x Rs. 500)")]
    steps = ["OPEN | /cart", "ASSERT_NUMBER_EQUALS | css=.qty | 4",
             "ASSERT_MULTIPLY | css=.price | css=.qty | css=.total",
             "ASSERT_NUMBER_EQUALS | css=.total | 2000"]
    good = validate_conversion(source, _proposal("TC_014", steps, steps[1:]), "https://example.test")
    assert good[0]["conversion_note"] == ""
    weak = validate_conversion(
        source, _proposal("TC_014", steps[:-1], steps[1:-1]), "https://example.test",
    )
    assert "cart total" in weak[0]["conversion_note"].casefold()


def test_timing_without_threshold_remains_reviewed():
    source = [_case("TC_001", "Home loads within acceptable time")]
    proposal = _proposal("TC_001", ["OPEN | /", "ASSERT_URL | /"], ["ASSERT_URL | /"])
    converted = validate_conversion(source, proposal, "https://example.test")
    assert "numeric time limit" in converted[0]["conversion_note"]
    invented = _proposal("TC_001", ["OPEN | /", "ASSERT_LOAD_TIME | 3000"],
                         ["ASSERT_LOAD_TIME | 3000"])
    converted = validate_conversion(source, invented, "https://example.test")
    assert not converted[0]["execute_partial"]


def test_conversion_rejects_synthetic_metadata_and_late_alert_watch():
    source = [_case("TC_009", "'Done' is displayed")]
    source[0]["step"] = "Click 'Submit' and accept the browser alert"
    steps = ["OPEN | /", "CLICK | button[data-form-action='/contact']",
             "WATCH_DIALOG | accept", "ASSERT_DIALOG | any", "ASSERT_TEXT | css=#done | Done"]
    result = validate_conversion(
        source, _proposal("TC_009", steps, [steps[-1]]), "https://example.test",
        allow_site_actions=True,
    )
    assert "snapshot-only" in result[0]["conversion_note"] or "alert" in result[0]["conversion_note"]
    steps[1] = "CLICK | css=#submit"
    result = validate_conversion(
        source, _proposal("TC_009", steps, [steps[-1]]), "https://example.test",
        allow_site_actions=True,
    )
    assert "alert" in result[0]["conversion_note"]


def test_account_created_must_be_checked_before_continue():
    source = [_case("TC_002", "'ACCOUNT CREATED!' is displayed; after Continue, header shows 'Logged in as QA Tester'")]
    steps = ["OPEN | /register", "CLICK | Continue",
             "ASSERT_TEXT | css=h2 | ACCOUNT CREATED!",
             "ASSERT_TEXT | css=.header | Logged in as QA Tester"]
    proposal = {"cases": [{"id": "TC_002", "steps": steps, "checks": [
        {"clause_index": 0, "assertions": [steps[2]]},
        {"clause_index": 1, "assertions": [steps[3]]},
    ]}]}
    converted = validate_conversion(source, proposal, "https://example.test")
    assert "before clicking Continue" in converted[0]["conversion_note"]


def test_empty_cart_does_not_pass_without_adding_and_removing_product():
    source = [_case("TC_015", "Product is removed from the cart; when empty, 'Cart is empty!' message is displayed")]
    source[0]["step"] = "Add a product to cart\nClick 'Cart'\nClick 'X' (delete) icon against the product"
    proposal = {"cases": [{"id": "TC_015", "steps": [
        "OPEN | /view_cart", "ASSERT_COUNT | css=.cart-product | 0",
        "ASSERT_TEXT | css=.cart-empty | Cart is empty!",
    ], "checks": [
        {"clause_index": 0, "assertions": ["ASSERT_COUNT | css=.cart-product | 0"]},
        {"clause_index": 1, "assertions": ["ASSERT_TEXT | css=.cart-empty | Cart is empty!"]},
    ]}]}
    checked = validate_conversion(source, proposal, "https://example.test", allow_site_actions=True)
    assert "click action" in checked[0]["conversion_note"]
    assert not checked[0]["execute_partial"]


def test_dom_modal_cannot_be_asserted_as_a_browser_dialog():
    source = [_case("TC_016", "Modal displays 'Register / Login account to proceed on checkout.'")]
    steps = ["OPEN | /cart", "ASSERT_DIALOG | Register / Login account to proceed on checkout."]
    checked = validate_conversion(
        source, _proposal("TC_016", steps, [steps[-1]]), "https://example.test",
    )
    assert "modal" in checked[0]["conversion_note"].casefold()


def test_retry_replaces_only_reviewed_rows(tmp_path: Path):
    source = tmp_path / "source.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append(["TestCase ID", "Title", "Test Step", "Test Data", "Expected Result"])
    sheet.append(["TC_001", "Good", "Open home", "", "'Done' title is displayed"])
    sheet.append(["TC_002", "Review", "Open home", "", "'Done' title is displayed"])
    book.save(source)
    originals = [_case("TC_001", "'Done' title is displayed"), _case("TC_002", "'Done' title is displayed")]
    good = validate_conversion(
        originals[:1], _proposal("TC_001", ["OPEN | /", "ASSERT_TITLE | Done"],
                                 ["ASSERT_TITLE | Done"]), "https://example.test",
    )[0]
    blocked = validate_conversion(originals[1:], {"cases": [{"id": "TC_002", "steps": [],
                                                              "checks": [], "review_reason": "Needs repair"}]},
                                  "https://example.test")[0]
    path = converted_suite_path(tmp_path, "https://example.test", source)
    save_converted_suite(source, path, [good, blocked])
    assert reviewed_case_ids(path) == ["TC_002"]
    replacement = validate_conversion(
        originals, {"cases": [
            {"id": "TC_001", "steps": [], "checks": [], "review_reason": "bad proposal"},
            {"id": "TC_002", "steps": ["OPEN | /done", "ASSERT_TITLE | Done"],
             "checks": [{"clause_index": 0, "assertions": ["ASSERT_TITLE | Done"]}],
             "review_reason": ""},
        ]}, "https://example.test",
    )
    assert replace_reviewed_conversions(path, replacement) == 1
    assert reviewed_case_ids(path) == []
    from selfheal.suite_loader import read_excel_tests

    rows = read_excel_tests(path)
    assert rows[0]["step"] == good["step"]
    assert rows[1]["step"] == replacement[1]["step"]


def test_generic_assertions_check_real_page_state(tmp_path: Path):
    config = {"environments": {"prod": {"base_url": "https://example.test"}}}
    executor = GenericStepExecutor(config, "prod", tmp_path / "suite.xlsx", 5000)
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=True, channel="chromium")
        try:
            page = browser.new_page()
            page.route("https://example.test/", lambda route: route.fulfill(
                status=200, content_type="text/html", body="""
                <h1>Products</h1><div class='name'>Top One</div><div class='name'>Top Two</div>
                <span id='price'>Rs. 500</span><input id='qty' value='4'>
                <span id='total'>Rs. 2,000</span><span id='count'>2 products</span>
                <div id='address'>Main Street, City, 12345</div>
                <button onclick="alert('Confirmed')">Submit</button>
                """))
            assert executor.execute(page, "OPEN | /", {})
            assert executor.execute(page, "ASSERT_LOAD_TIME | 5000", {})
            assert executor.execute(page, "ASSERT_COUNT | css=.name | 2", {})
            assert executor.execute(page, "ASSERT_EACH_TEXT_CONTAINS | css=.name | Top", {})
            assert executor.execute(page, "ASSERT_COUNT_MATCHES | css=.name | css=#count", {})
            assert executor.execute(page, "ASSERT_MULTIPLY | css=#price | css=#qty | css=#total", {})
            assert executor.execute(page, "ASSERT_NUMBER_EQUALS | css=#total | 2000", {})
            assert executor.execute(page, "ASSERT_DATA_FIELDS | css=#address | address,city,zip",
                                    {"address": "Main Street", "city": "City", "zip": "12345"})
            executor.execute(page, "WATCH_DIALOG | accept", {})
            executor.execute(page, "CLICK | css=button", {})
            assert executor.execute(page, "ASSERT_DIALOG | Confirmed", {})
            assert executor.execute(page, "ASSERT_URL_NOT_CONTAINS | checkout", {})
            with pytest.raises(Exception, match="Count mismatch"):
                executor.execute(page, "ASSERT_COUNT | css=.name | 0", {})
        finally:
            browser.close()


def test_missing_direct_text_locator_can_heal_without_weakening_the_assertion(tmp_path: Path):
    class Healer:
        def __init__(self):
            self.ai_calls = 0

        def resolve(self, page, target, *, kind, allow_ai=True, expected_text=None):
            if not allow_ai:
                return None
            self.ai_calls += 1
            if expected_text == "Registration failed" and page.locator("#actual").count() == 1:
                return HealingMatch(page.locator("#actual"), "healed")
            return None

        def remember_success(self, _match):
            pass

    healer = Healer()
    config = {
        "environments": {"prod": {"base_url": "https://example.test"}},
        "resolver": {"discovery_mode": "batch_cached"},
    }
    executor = GenericStepExecutor(config, "prod", tmp_path / "suite.xlsx", 5000, healer=healer)
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=True, channel="chromium")
        try:
            page = browser.new_page()
            page.set_content("<div id='actual'>Registration failed</div>")
            assert executor.execute(page, "ASSERT_TEXT | css=.stale | Registration failed", {})
            assert healer.ai_calls == 1
            with pytest.raises(Exception, match="Visible locator not found"):
                executor.execute(page, "ASSERT_TEXT | css=.stale | Not present", {})
        finally:
            browser.close()


def test_visible_assertion_does_not_accept_partial_target_words():
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=True, channel="chromium")
        try:
            page = browser.new_page()
            page.set_content("<h2>Login to your account</h2><a href='/login'>Signup / Login</a>")
            assert verified_locator(page, "h2", "ASSERT_VISIBLE", "Signup / Login") is None
            assert verified_locator(page, "a[href='/login']", "ASSERT_VISIBLE", "Signup / Login")
        finally:
            browser.close()
