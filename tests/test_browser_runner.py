"""Browser-level checks for actions and honest suite results."""

import json

import openpyxl
import pytest
import yaml
from openpyxl import Workbook
from playwright.sync_api import sync_playwright
from typer.testing import CliRunner

from selfheal.cli import app, run_browser_suite
from selfheal.executor.browser import ActionFailedError, BrowserExecutor

LOGIN_PAGE = """<!doctype html>
<html><body>
<a href="#auth" id="open-auth">Signup / Login</a>
<div id="auth" hidden>
  <section>
    <h2>New User Signup!</h2>
    <form id="signup"><input name="name" placeholder="Name">
      <input name="email" type="email" placeholder="Email Address">
      <button>Signup</button></form>
  </section>
  <section>
    <h2>Login to your account</h2>
    <form id="login"><input name="email" type="email" placeholder="Email Address" required>
      <input name="password" type="password" placeholder="Password" required>
      <button>Login</button></form>
  </section>
</div>
<div id="output"></div>
<script>
document.querySelector('#open-auth').onclick = event => {
  event.preventDefault(); document.querySelector('#auth').hidden = false;
};
document.querySelector('#login').onsubmit = event => {
  event.preventDefault();
  const form = event.target;
  const output = document.querySelector('#output');
  if (form.email.value === 'registered@example.test' && form.password.value === 'correct-password') {
    output.innerHTML = '<span>Logged in as QA Tester</span><a href="#logout">Logout</a>' +
      '<a href="#delete">Delete Account</a>';
  } else {
    output.textContent = 'Your email or password is incorrect!';
  }
};
</script>
</body></html>"""

REGISTRATION_PAGE = """<!doctype html><html><body>
<a href="#auth" id="open-auth">Signup / Login</a>
<section id="auth" hidden>
  <h2>New User Signup!</h2>
  <form id="signup"><input name="name" placeholder="Name">
    <input name="email" type="email" placeholder="Email Address">
    <button>Signup</button></form>
</section>
<section id="register" hidden>
  <form id="account">
    <label><input type="radio" name="title" value="Mr">Mr.</label>
    <label><input type="radio" name="title" value="Mrs">Mrs.</label>
    <input type="password" name="password" placeholder="Password">
    <select name="days"><option value="10">10</option></select>
    <select name="months"><option value="5">May</option></select>
    <select name="years"><option value="1995">1995</option></select>
    <input name="first_name" placeholder="First name">
    <input name="last_name" placeholder="Last name">
    <input name="address1" placeholder="Address">
    <input name="address2" placeholder="Address 2">
    <select name="country"><option value="India">India</option></select>
    <input name="state" placeholder="State">
    <input name="city" placeholder="City">
    <input name="zipcode" placeholder="Zipcode">
    <input name="mobile_number" placeholder="Mobile Number">
    <button>Create Account</button>
  </form>
</section>
<section id="created" hidden><h1>ACCOUNT CREATED!</h1><button id="continue">Continue</button></section>
<div id="output"></div>
<script>
document.querySelector('#open-auth').onclick = event => {
  event.preventDefault(); document.querySelector('#auth').hidden = false;
};
document.querySelector('#signup').onsubmit = event => {
  event.preventDefault(); document.querySelector('#auth').hidden = true;
  document.querySelector('#register').hidden = false;
};
document.querySelector('#account').onsubmit = event => {
  event.preventDefault(); document.querySelector('#register').hidden = true;
  document.querySelector('#created').hidden = false;
};
document.querySelector('#continue').onclick = () => {
  document.querySelector('#created').hidden = true;
  document.querySelector('#output').textContent = 'Logged in as QA Tester';
};
</script>
</body></html>"""


def config_for(url, *, ai=False, auto_approve=False):
    return {
        "app_name": "fixture",
        "environments": {
            "prod": {
                "base_url": url,
                "account": {"email": "registered@example.test", "password": "correct-password"},
            }
        },
        "execution": {"browser": "chromium", "timeout_seconds": 5},
        "resolver": {"ai_enabled": ai, "confidence_threshold": 0.85},
        "healing": {"enabled": ai, "auto_approve": auto_approve, "max_heals_per_run": 5},
        "safety": {"no_real_payments": True},
    }


@pytest.fixture
def browser():
    with sync_playwright() as playwright:
        instance = playwright.chromium.launch(headless=True, channel="chromium")
        yield instance
        instance.close()


@pytest.fixture
def login_page(tmp_path, monkeypatch):
    monkeypatch.setenv("TEST_EMAIL", "registered@example.test")
    monkeypatch.setenv("TEST_PASSWORD", "correct-password")
    path = tmp_path / "login.html"
    path.write_text(LOGIN_PAGE, encoding="utf-8")
    return path


def execute(browser, page_file, case, *, config=None, executor=None):
    context = browser.new_context()
    try:
        page = context.new_page()
        runner = executor or BrowserExecutor(config or config_for(page_file.as_uri()), "prod", page_file)
        return runner.run_case(page, case)
    finally:
        context.close()


def test_login_clicks_and_verifies_real_state(browser, login_page):
    case = {
        "step": "1. Click 'Signup / Login'\n2. Enter registered email and correct password "
                "under 'Login to your account'\n3. Click 'Login'",
        "data": "Email: <registered email>\nPassword: <registered password>",
        "expected": "User is logged in; header shows 'Logged in as <username>' "
                    "along with 'Logout' and 'Delete Account' links.",
    }
    result = execute(browser, login_page, case)
    assert result["status"] == "passed", result
    assert [step["status"] for step in result["steps"]] == ["passed"] * 3
    assert "Clicked Login" in result["steps"][2]["detail"]


def test_invalid_and_blank_login_are_verified(browser, login_page):
    invalid = {
        "step": "Click 'Signup / Login'\nEnter registered email and wrong password\nClick 'Login'",
        "data": "Email: <registered email>\nPassword: wrong-password",
        "expected": "Error 'Your email or password is incorrect!' is displayed; "
                    "user is not logged in.",
    }
    blank = {
        "step": "Click 'Signup / Login'\nLeave Email and Password empty\nClick 'Login'",
        "data": "Email: (blank)\nPassword: (blank)",
        "expected": "HTML5 required-field validation ('Please fill out this field') "
                    "appears on Email; form is not submitted.",
    }
    assert execute(browser, login_page, invalid)["status"] == "passed"
    assert execute(browser, login_page, blank)["status"] == "passed"


def test_workbook_registration_steps_execute_before_review(browser, tmp_path):
    page_file = tmp_path / "registration.html"
    page_file.write_text(REGISTRATION_PAGE, encoding="utf-8")
    from selfheal.cli import read_excel_tests

    case = next(item for item in read_excel_tests("suites/tests.xlsx") if item["id"] == "TC_002")
    result = execute(browser, page_file, case)
    assert result["status"] == "needs_review", result
    assert len(result["steps"]) == 7
    assert all(item["status"] == "passed" for item in result["steps"])
    assert all(item["status"] == "passed" for item in result["checks"])
    assert any("Title" in item for item in result["notes"])

    configured = config_for(page_file.as_uri())
    configured["registration"] = {"default_title": "Mr"}
    result = execute(browser, page_file, case, config=configured)
    assert result["status"] == "passed", result
    assert len(result["steps"]) == 7


def test_contact_upload_and_alert_are_real_actions(browser, tmp_path):
    page_file = tmp_path / "contact.html"
    page_file.write_text(
        "<a href='#home'>Home</a>"
        "<a href='#contact' onclick=\"event.preventDefault();document.querySelector('#contact').hidden=false\">"
        "Contact us</a><section id='contact' hidden><form id='contact-form'>"
        "<input name='name' placeholder='Name'><input type='email' name='email' placeholder='Email'>"
        "<input name='subject' placeholder='Subject'><textarea name='message'></textarea>"
        "<input type='file' name='attachment'><button>Submit</button></form></section>"
        "<p id='out'></p><a href='#home' class='btn' id='success-home' hidden>Home</a>"
        "<script>document.querySelector('#contact-form').onsubmit=event=>{"
        "event.preventDefault();alert('Confirm');document.querySelector('#out').textContent="
        "'Success! Your details have been submitted successfully.';"
        "document.querySelector('#success-home').hidden=false};"
        "document.querySelector('#success-home').onclick=event=>{"
        "event.preventDefault();document.querySelector('#out').textContent='Returned home';"
        "document.querySelector('#contact').hidden=true}</script>",
        encoding="utf-8",
    )
    executor = BrowserExecutor(config_for(page_file.as_uri()), "prod", tmp_path / "tests.xlsx")
    (tmp_path / "sample.txt").write_text("contact fixture", encoding="utf-8")
    result = execute(
        browser, page_file,
        {"step": "Click 'Contact us'\nEnter Name, Email, Subject, Message\n"
                 "Upload a file\nClick 'Submit' and accept the browser alert\nClick 'Home'",
         "data": "Name: QA Tester\nEmail: qa@example.test\nSubject: X\n"
                 "Message: Y\nFile: sample.txt",
         "expected": "'Success! Your details have been submitted successfully.' is displayed; "
                     "'Home' button returns user to home page; 'Returned home' is displayed"},
        executor=executor,
    )
    assert result["status"] == "passed", result
    assert [step["status"] for step in result["steps"]] == ["passed"] * 5


def test_skipped_or_incorrect_steps_never_pass(browser, login_page):
    unsupported = {
        "step": "Click 'Signup / Login'\nTeleport to dashboard",
        "expected": "'Logged in as QA Tester' is displayed",
    }
    missing = {"step": "Click 'Does not exist'", "expected": "'Done' is displayed"}
    wrong_assertion = {
        "step": "Click 'Signup / Login'\nEnter registered email and wrong password\nClick 'Login'",
        "data": "Email: <registered email>\nPassword: wrong-password",
        "expected": "'Logged in as <username>' is displayed",
    }
    assert execute(browser, login_page, unsupported)["status"] == "needs_review"
    assert execute(browser, login_page, missing)["status"] == "failed"
    assert execute(browser, login_page, wrong_assertion)["status"] == "failed"


def test_generic_commands_run_a_configured_site_without_site_specific_code(browser, tmp_path):
    page_file = tmp_path / "generic_portal.html"
    page_file.write_text(
        "<title>Generic Portal</title><form id='login'>"
        "<label>Work email <input id='email' type='email'></label>"
        "<select id='plan'><option>Basic</option><option>Premium</option></select>"
        "<label><input data-testid='terms' type='checkbox'> Terms</label>"
        "<button type='submit'>Sign in</button></form>"
        "<section class='dashboard' hidden></section>"
        "<script>document.querySelector('#login').onsubmit=event=>{event.preventDefault();"
        "document.querySelector('.dashboard').hidden=false;"
        "document.querySelector('.dashboard').textContent='Welcome '+document.querySelector('#email').value;"
        "location.hash='dashboard'}</script>",
        encoding="utf-8",
    )
    config = config_for(page_file.as_uri())
    config["locators"] = {
        "email": {"by": "label", "value": "Work email"},
        "plan": "css=#plan",
        "terms": "testid=terms",
        "sign_in": {"by": "role", "role": "button", "name": "Sign in"},
        "dashboard": "css=.dashboard",
    }
    result = execute(
        browser,
        page_file,
        {
            "step": "\n".join(
                (
                    "OPEN | ${BASE_URL}",
                    "FILL | email | ${DATA:email}",
                    "ASSERT_VALUE | email | ${DATA:email}",
                    "SELECT | plan | Premium",
                    "CHECK | terms",
                    "CLICK | sign_in",
                    "WAIT_VISIBLE | dashboard",
                    "ASSERT_TEXT | dashboard | Welcome qa@example.test",
                    "ASSERT_URL_CONTAINS | #dashboard",
                    "ASSERT_TITLE | Generic Portal",
                )
            ),
            "data": "Email: qa@example.test",
            "expected": "",
        },
        config=config,
    )
    assert result["status"] == "passed", result
    assert len(result["steps"]) == 10
    assert result["checks"] == [{"status": "passed", "expectation": "Inline generic assertions", "reason": None}]


def test_icon_font_prefix_does_not_hide_a_real_link(browser, tmp_path):
    page_file = tmp_path / "products.html"
    page_file.write_text(
        "<a href='#products' onclick=\"document.querySelector('#out').textContent='ALL PRODUCTS'\">"
        "<i>&#xe8f8;</i> Products</a><h1 id='out'></h1>",
        encoding="utf-8",
    )
    result = execute(
        browser, page_file,
        {"step": "Click 'Products'", "expected": "'ALL PRODUCTS' is displayed"},
    )
    assert result["status"] == "passed", result


def test_brand_count_prefix_does_not_hide_link_name(browser, tmp_path):
    page_file = tmp_path / "brand.html"
    page_file.write_text(
        "<a href='#brand' onclick=\"document.querySelector('#out').textContent="
        "'BRAND - POLO PRODUCTS'\">(6)Polo</a><h1 id='out'></h1>",
        encoding="utf-8",
    )
    result = execute(
        browser, page_file,
        {"step": "Click brand 'Polo' in Brands sidebar",
         "expected": "'BRAND - POLO PRODUCTS' is displayed"},
    )
    assert result["status"] == "passed", result


def test_review_uses_first_product_when_workbook_does_not_name_one(browser, tmp_path):
    page_file = tmp_path / "reviews.html"
    page_file.write_text(
        "<a href='#products' onclick=\"document.querySelector('#listing').hidden=false\">Products</a>"
        "<section id='listing' hidden>"
        "<div class='product-image-wrapper'><div class='productinfo'><p>Blue Top</p></div>"
        "<a href='#first' onclick=\"document.querySelector('#detail').hidden=false;"
        "document.querySelector('#listing').hidden=true\">View Product</a></div>"
        "<div class='product-image-wrapper'><div class='productinfo'><p>Other Top</p></div>"
        "<a href='#second'>View Product</a></div></section>"
        "<section id='detail' hidden><h2>Write Your Review</h2>"
        "<form id='review'><input name='name' placeholder='Your Name'>"
        "<input name='email' placeholder='Email Address'>"
        "<textarea name='review' placeholder='Add Review Here!'></textarea>"
        "<button>Submit</button></form><p id='out'></p></section>"
        "<script>document.querySelector('#review').onsubmit=event=>{"
        "event.preventDefault();"
        "document.querySelector('#out').textContent='Thank you for your review.'}</script>",
        encoding="utf-8",
    )
    result = execute(
        browser, page_file,
        {"step": "Click 'Products' > 'View Product'\nLocate 'Write Your Review'\n"
                 "Enter name, email and review\nClick 'Submit'",
         "data": "Name: QA Tester\nEmail: qa@example.test\nReview: Good product",
         "expected": "Message 'Thank you for your review.' is displayed."},
    )
    assert result["status"] == "passed", result
    assert len(result["steps"]) == 4


def test_search_icon_next_to_input_works_without_form(browser, tmp_path):
    page_file = tmp_path / "search.html"
    page_file.write_text(
        "<div><input placeholder='Search Product'>"
        "<button onclick=\"document.querySelector('#out').textContent='SEARCHED PRODUCTS'\">"
        "<i>&#xe8f8;</i></button></div><h1 id='out'></h1>",
        encoding="utf-8",
    )
    result = execute(
        browser, page_file,
        {"step": "Enter keyword in search box\nClick search icon",
         "data": "Search: Top", "expected": "'SEARCHED PRODUCTS' is displayed"},
    )
    assert result["status"] == "passed", result


def test_search_recovers_once_from_unexpected_navigation(browser, tmp_path):
    page_file = tmp_path / "search-retry.html"
    (tmp_path / "wrong-page.html").write_text("<h1>Wrong page</h1>", encoding="utf-8")
    page_file.write_text(
        "<input id='search_product' name='search' placeholder='Search Product'>"
        "<button id='submit_search' onclick=\"if(!window.name){window.name='clicked';"
        "location.href='wrong-page.html'}else{document.querySelector('#out').textContent="
        "'SEARCHED PRODUCTS'}\">Search</button><h2 id='out'></h2>",
        encoding="utf-8",
    )
    config = config_for(page_file.as_uri())
    config["execution"].update({"retry_on_failure": True, "max_retries_per_step": 1})
    result = execute(
        browser, page_file,
        {"step": "Enter keyword in search box\nClick search icon",
         "data": "Search: Top", "expected": "'SEARCHED PRODUCTS' is displayed"},
        config=config,
    )
    assert result["status"] == "passed", result
    assert "retried after unexpected navigation" in result["steps"][1]["detail"]


def test_brand_count_can_be_read_on_result_page(browser, tmp_path):
    page_file = tmp_path / "brand-count.html"
    page_file.write_text(
        "<div class='brands-name'><a href='/brand_products/Polo'>"
        "<span>(2)</span>Polo</a></div>",
        encoding="utf-8",
    )
    config = config_for(page_file.as_uri())
    executor = BrowserExecutor(config, "prod", page_file)
    executor.listings = {
        "BRAND - POLO PRODUCTS": [
            {"name": "Blue Top", "href": "/product_details/1"},
            {"name": "Fancy Green Top", "href": "/product_details/8"},
        ]
    }
    executor._detail_for_card = lambda _page, card: {"name": card["name"], "brand": "Polo"}
    context = browser.new_context()
    try:
        page = context.new_page()
        page.goto(page_file.as_uri())
        executor._verify_filtered_listing(
            page, "Brand page titled 'BRAND - POLO PRODUCTS' shows only Polo products "
                  "(count matches sidebar)", {"brand": "Polo"}
        )
    finally:
        context.close()


def test_search_checks_all_names_or_visible_product_categories(tmp_path):
    executor = BrowserExecutor(config_for((tmp_path / "search.html").as_uri()), "prod", tmp_path)
    executor.listings = {
        "SEARCHED PRODUCTS": [
            {"name": "Blue Top", "href": "/product_details/1"},
            {"name": "Little Girls Mr. Panda Shirt", "href": "/product_details/18"},
        ]
    }
    data = {"search": "Top"}
    with pytest.raises(ActionFailedError, match="does not contain"):
        executor._verify_search_results(
            None, "'SEARCHED PRODUCTS' heading is shown and every displayed product name "
                  "contains the keyword 'Top'", data
        )
    executor._detail_for_card = lambda _page, card: {
        "name": card["name"], "category": "Kids > Tops & Shirts"
    }
    executor._verify_search_results(
        None, "'SEARCHED PRODUCTS' heading is shown and every displayed product name or "
              "category contains the keyword 'Top'", data
    )


def test_filtered_products_check_detail_metadata_and_sidebar_count(tmp_path):
    executor = BrowserExecutor(config_for((tmp_path / "filters.html").as_uri()), "prod", tmp_path)
    executor.listings = {
        "WOMEN - TOPS PRODUCTS": [{"name": "Blue Top", "href": "/product_details/1"}],
        "BRAND - POLO PRODUCTS": [{"name": "Blue Top", "href": "/product_details/1"}],
    }
    executor.brand_sidebar_count = 1
    executor._detail_for_card = lambda _page, _card: {
        "name": "Blue Top", "category": "Women > Tops", "brand": "Polo"
    }
    data = {"category": "Women > Tops", "brand": "Polo"}
    executor._verify_filtered_listing(
        None, "Page titled 'WOMEN - TOPS PRODUCTS' shows only women tops", data
    )
    executor._verify_filtered_listing(
        None, "Brand page titled 'BRAND - POLO PRODUCTS' shows only Polo products (count matches sidebar)", data
    )
    executor.brand_sidebar_count = 2
    with pytest.raises(ActionFailedError, match="count differs"):
        executor._verify_filtered_listing(
            None, "Brand page titled 'BRAND - POLO PRODUCTS' shows only Polo products (count matches sidebar)", data
        )


def test_invalid_subscription_is_blocked_even_with_prior_success_banner(browser, tmp_path):
    page_file = tmp_path / "subscribe.html"
    page_file.write_text(
        "<footer><form><input name='email' type='email' required>"
        "<button type='submit'>Go</button></form><p id='status'></p></footer>"
        "<script>document.querySelector('form').onsubmit=e=>{e.preventDefault();"
        "document.querySelector('#status').textContent="
        "'You have been successfully subscribed!'}</script>",
        encoding="utf-8",
    )
    result = execute(
        browser, page_file,
        {"step": "Scroll to footer on home page\nEnter email in Subscription field\n"
                 "Click arrow button\nRepeat with invalid email 'abc'",
         "data": "Valid: qa@example.test\nInvalid: abc",
         "expected": "Valid email shows 'You have been successfully subscribed!'; "
                     "invalid email triggers email format validation and submission is blocked."},
    )
    assert result["status"] == "passed", result
    assert len(result["steps"]) == 4


def test_enter_comment_happens_before_click(browser, tmp_path):
    page_file = tmp_path / "comment.html"
    page_file.write_text(
        "<textarea name='message'></textarea><button onclick=\"document.querySelector('#out')."
        "textContent=document.querySelector('textarea').value\">Place Order</button>"
        "<p id='out'></p>",
        encoding="utf-8",
    )
    result = execute(
        browser, page_file,
        {"step": "Enter comment, click 'Place Order'",
         "data": "Comment: Deliver after 6 PM",
         "expected": "'Deliver after 6 PM' is displayed"},
    )
    assert result["status"] == "passed", result
    assert "Filled comment" in result["steps"][0]["detail"]


def test_demo_payment_requires_configured_host_and_card(browser, tmp_path):
    page_file = tmp_path / "payment.html"
    page_file.write_text(
        "<form><input name='name_on_card'><input name='card_number'>"
        "<input name='cvc'><input name='expiry_month'><input name='expiry_year'>"
        "<button>Pay and Confirm Order</button></form><p id='out'></p>"
        "<script>document.querySelector('form').onsubmit=e=>{e.preventDefault();"
        "document.querySelector('#out').textContent='ORDER PLACED!'}</script>",
        encoding="utf-8",
    )
    config = config_for(page_file.as_uri())
    config["safety"]["demo_orders"] = {
        "enabled": True, "hosts": [None], "card_numbers": ["4111111111111111"]
    }
    case = {
        "step": "Enter card details, click 'Pay and Confirm Order'",
        "data": "Name on Card: QA Tester\nCard No: 4111111111111111\nCVC: 123\nExpiry: 12/2028",
        "expected": "'ORDER PLACED!' is displayed",
    }
    assert execute(browser, page_file, case, config=config)["status"] == "passed"
    bad_card = dict(case, data=case["data"].replace("4111111111111111", "4000000000000000"))
    blocked = execute(browser, page_file, bad_card, config=config)
    assert blocked["status"] == "needs_review"
    assert "limited to the configured demo host and test card" in blocked["error"]
    config["safety"]["demo_orders"]["hosts"] = ["example.test"]
    blocked_host = execute(browser, page_file, case, config=config)
    assert blocked_host["status"] == "needs_review"


def test_multiple_cart_products_verify_unit_price_quantity_and_total(browser, tmp_path):
    page_file = tmp_path / "multi-cart.html"
    page_file.write_text(
        "<table><thead><tr><td>Item</td><td>Description</td><td>Price</td>"
        "<td>Quantity</td><td>Total</td></tr></thead><tbody>"
        "<tr><td></td><td>Blue Top</td><td>Rs. 500</td><td>1</td><td>Rs. 500</td></tr>"
        "<tr><td></td><td>Men Tshirt</td><td>Rs. 400</td><td>1</td><td>Rs. 400</td></tr>"
        "</tbody></table>",
        encoding="utf-8",
    )
    executor = BrowserExecutor(config_for(page_file.as_uri()), "prod", page_file)
    executor.add_feedback = ["Added!", "Added!"]
    executor.cart_products = ["Blue Top", "Men Tshirt"]
    data = {"product 1": "Blue Top (Rs. 500)", "product 2": "Men Tshirt (Rs. 400)"}
    context = browser.new_context()
    try:
        page = context.new_page()
        page.goto(page_file.as_uri())
        executor._verify_added_cart(page, "Cart shows both products with correct price", data)
        with pytest.raises(ActionFailedError, match="Cart price is wrong"):
            executor._verify_added_cart(
                page, "Cart shows both products with correct price",
                dict(data, **{"product 2": "Men Tshirt (Rs. 450)"}),
            )
    finally:
        context.close()


def test_delayed_modal_button_is_clicked(browser, tmp_path):
    page_file = tmp_path / "modal.html"
    page_file.write_text(
        "<button onclick=\"setTimeout(()=>document.querySelector('#modal').hidden=false,300)\">"
        "Add to cart</button><section id='modal' hidden><h1>Added!</h1><button onclick=\""
        "document.querySelector('#modal').hidden=true;"
        "document.querySelector('#out').textContent='Done'\">Continue Shopping</button>"
        "</section><p id='out'></p>",
        encoding="utf-8",
    )
    result = execute(
        browser, page_file,
        {"step": "Click 'Add to cart'\nClick 'Continue Shopping'",
         "expected": "'Added!' modal appears; 'Done' is displayed"},
    )
    assert result["status"] == "passed", result


def test_product_lookup_waits_for_named_product(browser, tmp_path):
    page_file = tmp_path / "product.html"
    page_file.write_text(
        "<div id='items'></div><p id='out'></p><script>setTimeout(()=>{"
        "document.querySelector('#items').innerHTML=\"<article><p>Blue Top</p>"
        "<button onclick='document.querySelector(\\\"#out\\\").textContent=\\\"Added!\\\"'>"
        "Add to cart</button></article>\"},300)</script>",
        encoding="utf-8",
    )
    result = execute(
        browser, page_file,
        {"step": "Add a product to cart", "data": "Product: Blue Top",
         "expected": "'Added!' is displayed"},
    )
    assert result["status"] == "passed", result


def test_delayed_add_reaches_cart_and_row_control_removes_product(browser, tmp_path):
    page_file = tmp_path / "cart.html"
    page_file.write_text(
        "<a id='cart-link' href='#cart'>Cart</a>"
        "<article id='listing'><p>Blue Top</p><button id='add'>Add to cart</button></article>"
        "<section class='modal' id='modal' hidden><h2>Added!</h2>"
        "<a id='view-cart' href='#cart'>View Cart</a></section>"
        "<table id='cart' hidden></table><p id='empty'></p>"
        "<script>let added=false;"
        "document.querySelector('#add').onclick=()=>setTimeout(()=>{"
        "added=true;document.querySelector('#modal').hidden=false},300);"
        "function showCart(event){event.preventDefault();"
        "document.querySelector('#listing').hidden=true;"
        "document.querySelector('#modal').hidden=true;"
        "const cart=document.querySelector('#cart');cart.hidden=false;"
        "cart.innerHTML=added?'<tr><td>Blue Top</td><td>"
        "<a href=\"#delete\" aria-label=\"Remove item\">×</a></td></tr>':'';"
        "document.querySelector('#empty').textContent=added?'':'Cart is empty!';}"
        "document.querySelector('#cart-link').onclick=showCart;"
        "document.querySelector('#view-cart').onclick=showCart;"
        "document.querySelector('#cart').onclick=event=>{"
        "if(event.target.closest('a')){event.preventDefault();"
        "event.target.closest('tr').remove();"
        "document.querySelector('#empty').textContent='Cart is empty!'}}"
        "</script>",
        encoding="utf-8",
    )
    result = execute(
        browser, page_file,
        {"step": "Add a product to cart\nClick 'Cart'\nClick 'X' (delete) icon against the product",
         "data": "Product: Blue Top",
         "expected": "Product is removed from the cart; when empty, 'Cart is empty!' message is displayed"},
    )
    assert result["status"] == "passed", json.dumps(result, indent=2)
    assert [item["status"] for item in result["steps"]] == ["passed"] * 3
    assert "(dialog)" in result["steps"][1]["detail"]


def test_clicking_add_without_cart_effect_fails(browser, tmp_path):
    page_file = tmp_path / "broken-cart.html"
    page_file.write_text(
        "<a href='#cart' onclick=\"document.querySelector('#listing').hidden=true;"
        "document.querySelector('#empty').textContent='Cart is empty!'\">Cart</a>"
        "<article id='listing'><p>Blue Top</p><button>Add to cart</button></article>"
        "<p id='empty'></p>",
        encoding="utf-8",
    )
    config = config_for(page_file.as_uri())
    config["execution"]["timeout_seconds"] = 2
    result = execute(
        browser, page_file,
        {"step": "Add a product to cart\nClick 'Cart'", "data": "Product: Blue Top",
         "expected": "'Blue Top' is shown in cart"},
        config=config,
    )
    assert result["status"] == "failed", result
    assert result["steps"][1]["status"] == "failed"
    assert "Cart does not show added product" in result["error"]


def test_guest_checkout_requires_visible_prompt_and_no_navigation(browser, tmp_path):
    page_file = tmp_path / "checkout.html"
    page_file.write_text(
        "<button onclick=\"document.querySelector('#guest').hidden=false\">"
        "Proceed To Checkout</button>"
        "<section id='guest' class='modal' hidden>"
        "<p>Register / Login account to proceed on checkout.</p>"
        "<a href='#login'>Register / Login</a></section>",
        encoding="utf-8",
    )
    result = execute(
        browser, page_file,
        {"step": "Click 'Proceed To Checkout'",
         "expected": "Modal displays 'Register / Login account to proceed on checkout.' "
                     "with a 'Register / Login' link; checkout page is not opened."},
    )
    assert result["status"] == "passed", result
    assert all(item["status"] == "passed" for item in result["checks"])


def test_cart_quantity_and_total_are_checked_against_table_cells(browser, tmp_path):
    page_file = tmp_path / "amounts.html"
    page_file.write_text(
        "<a href='#cart' onclick=\"document.querySelector('#cart').hidden=false\">Cart</a>"
        "<table id='cart' hidden><thead><tr><td>Item</td><td>Description</td>"
        "<td>Price</td><td>Quantity</td><td>Total</td></tr></thead>"
        "<tbody><tr><td></td><td>Blue Top</td><td>Rs. 500</td>"
        "<td>4</td><td>Rs. 2000</td></tr></tbody></table>",
        encoding="utf-8",
    )
    case = {
        "step": "Click 'Cart'",
        "data": "Product: Blue Top\nQuantity: 4",
        "expected": "Cart shows the product with quantity 4 and Total = Rs. 2000 (4 x Rs. 500).",
    }
    assert execute(browser, page_file, case)["status"] == "passed"
    case["expected"] = case["expected"].replace("Total = Rs. 2000", "Total = Rs. 2100")
    result = execute(browser, page_file, case)
    assert result["status"] == "failed", result
    assert "Cart total" in result["checks"][0]["reason"]


def test_only_named_account_placeholders_are_resolved(tmp_path, monkeypatch):
    monkeypatch.setenv("TEST_EMAIL", "registered@example.test")
    monkeypatch.setenv("TEST_PASSWORD", "correct-password")
    executor = BrowserExecutor(config_for((tmp_path / "login.html").as_uri()), "prod", tmp_path)
    data = executor._prepare_data(
        "Email: <registered email>\nPassword: <registered password>\nOther Password: <wrong password>"
    )
    assert data == {
        "email": "registered@example.test",
        "password": "correct-password",
        "other password": "<wrong password>",
    }
    unresolved = executor._prepare_data(
        "Email: <previously registered email>\nOther Email: <test account email>"
    )
    assert unresolved["email"] == "<previously registered email>"
    assert unresolved["other email"] == "<test account email>"
    executor.registered_account = {"email": "new@example.test", "password": "new-password"}
    registered = executor._prepare_data(
        "Email: <registered email>\nPassword: <registered password>\n"
        "Other Email: <test account email>\nTest Email: <test email>"
    )
    assert registered == {
        "email": "new@example.test",
        "password": "new-password",
        "other email": "new@example.test",
        "test email": "registered@example.test",
    }


def test_login_without_case_credentials_uses_configured_account(browser, login_page):
    executor = BrowserExecutor(config_for(login_page.as_uri()), "prod", login_page)
    executor.registered_account = {"email": "deleted@example.test", "password": "different"}
    result = execute(
        browser, login_page,
        {"step": "Login", "expected": "'Logged in as QA Tester' is displayed"},
        executor=executor,
    )
    assert result["status"] == "passed", result


def test_login_form_with_zero_height_wrapper_is_still_usable(browser, tmp_path, monkeypatch):
    monkeypatch.setenv("TEST_EMAIL", "registered@example.test")
    monkeypatch.setenv("TEST_PASSWORD", "correct-password")
    page_file = tmp_path / "flat-form.html"
    page_file.write_text(
        "<form style='display:contents'><input type='email' name='email'>"
        "<input type='password' name='password'><button>Login</button></form>"
        "<p id='out'></p><script>document.querySelector('form').onsubmit=e=>{"
        "e.preventDefault();document.querySelector('#out').textContent="
        "'Logged in as QA Tester'}</script>",
        encoding="utf-8",
    )
    result = execute(
        browser, page_file,
        {"step": "Login", "expected": "'Logged in as QA Tester' is displayed"},
    )
    assert result["status"] == "passed", result


def test_origin_error_page_is_reported_as_failure(browser, tmp_path):
    page_file = tmp_path / "origin-error.html"
    page_file.write_text(
        "<h1>Web server is returning an unknown error</h1><p>Error code 520</p>",
        encoding="utf-8",
    )
    result = execute(
        browser, page_file,
        {"step": "Observe the page", "expected": "'Done' is displayed"},
    )
    assert result["status"] == "failed"
    assert "HTTP 520" in result["error"]


def test_first_product_detail_uses_stable_link_href(browser, tmp_path):
    page_file = tmp_path / "first-link.html"
    page_file.write_text(
        "<div class='product-image-wrapper'><div class='productinfo'><p>Blue Top</p></div>"
        "<a href='#detail' onclick=\"document.querySelector('#out').textContent="
        "'Detail Opened'\"><i>&#xe8f8;</i> View Product</a></div><p id='out'></p>",
        encoding="utf-8",
    )
    result = execute(
        browser, page_file,
        {"step": "Click 'View Product' on the first product",
         "expected": "'Detail Opened' is displayed"},
    )
    assert result["status"] == "passed", result
    assert "product link" in result["steps"][0]["detail"]


class MemoryCache:
    def __init__(self):
        self.values = {}

    def get(self, app, environment, page, target):
        return self.values.get((app, environment, page, target))

    def put(self, app, environment, page, target, locator):
        self.values[(app, environment, page, target)] = locator


class FakeClient:
    def identify_element(self, page_tree, target, confidence_threshold):
        elements = json.loads(page_tree)
        button = next(item for item in elements if item["tag"] == "button")
        return {"found": True, "confidence": 0.99, "locator": "#" + button["id"]}


def test_ai_selector_is_clicked_and_stale_selector_is_healed(browser, tmp_path, monkeypatch):
    page_file = tmp_path / "action.html"
    config = config_for(page_file.as_uri(), ai=True, auto_approve=True)
    executor = BrowserExecutor(config, "prod", page_file)
    executor.healer.cache = MemoryCache()
    executor.healer.client = FakeClient()
    monkeypatch.setenv("OPENAI_API_KEY", "local-test-key")
    case = {"step": "Click 'Continue'", "expected": "'Done' is displayed"}
    for button_id, source in (("old-action", "ai"), ("new-action", "healed")):
        page_file.write_text(
            f"<button id='{button_id}' onclick=\"document.querySelector('#out').textContent='Done'\">Go</button>"
            "<p id='out'></p>",
            encoding="utf-8",
        )
        result = execute(browser, page_file, case, executor=executor)
        assert result["status"] == "passed", result
        assert f"({source})" in result["steps"][0]["detail"]


def test_ai_suggestion_requires_approval_when_configured(browser, tmp_path, monkeypatch):
    page_file = tmp_path / "action.html"
    page_file.write_text("<button id='action'>Go</button><p>Done</p>", encoding="utf-8")
    executor = BrowserExecutor(config_for(page_file.as_uri(), ai=True), "prod", page_file)
    executor.healer.cache = MemoryCache()
    executor.healer.client = FakeClient()
    monkeypatch.setenv("OPENAI_API_KEY", "local-test-key")
    result = execute(browser, page_file, {"step": "Click 'Continue'", "expected": "'Done' is displayed"}, executor=executor)
    assert result["status"] == "needs_review"
    assert "auto_approve is false" in result["error"]


def test_suite_results_match_browser_outcomes(login_page, tmp_path):
    config_file = tmp_path / "app.yaml"
    config_file.write_text(yaml.safe_dump(config_for(login_page.as_uri())), encoding="utf-8")
    cases = [
        {"id": "PASS", "title": "Login", "step": "Click 'Signup / Login'\n"
         "Enter registered email and correct password under 'Login to your account'\n"
         "Click 'Login'", "data": "Email: <registered email>\nPassword: <registered password>",
         "expected": "'Logged in as <username>' is displayed"},
        {"id": "REVIEW", "title": "Unknown", "step": "Do magic", "expected": "'Done' is displayed"},
        {"id": "FAIL", "title": "Missing", "step": "Click 'Missing'", "expected": "'Done' is displayed"},
    ]
    history_path = tmp_path / "run_history.xlsx"
    run_dir = run_browser_suite(
        cases,
        login_page,
        headless=True,
        config_path=config_file,
        output_root=tmp_path,
        history_path=history_path,
    )
    results = json.loads((run_dir / "results.json").read_text(encoding="utf-8"))
    assert (results["passed"], results["failed"], results["needs_review"]) == (1, 1, 1)
    assert results["status"] == "failed"
    assert [item["status"] for item in results["tests"]] == ["passed", "needs_review", "failed"]
    assert all(item.get("screenshot") for item in results["tests"])
    assert (run_dir / "screenshots" / "passed" / "001_PASS.png").is_file()
    assert (run_dir / "screenshots" / "review" / "002_REVIEW.png").is_file()
    assert (run_dir / "screenshots" / "failed" / "003_FAIL.png").is_file()
    history = openpyxl.load_workbook(history_path)["Runs"]
    assert history.max_row == 2
    assert [cell.value for cell in history[1]][:14] == [
        "Run ID", "Started At", "Finished At", "Duration (s)", "App", "Suite",
        "Environment", "Browser", "Headless", "Total", "Passed", "Failed",
        "Needs Review", "Status",
    ]
    assert [history.cell(row=2, column=column).value for column in (10, 11, 12, 13, 14)] == [3, 1, 1, 1, "failed"]


def test_cli_returns_failure_for_failed_case_and_saves_results(login_page, tmp_path, monkeypatch):
    config_file = tmp_path / "app.yaml"
    config_file.write_text(yaml.safe_dump(config_for(login_page.as_uri())), encoding="utf-8")
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["TestCase ID", "Title", "Test Step", "Test Data", "Expected Result", "Execute"])
    sheet.append(["TC_OK", "Invalid login", "Click 'Signup / Login'\n"
                  "Enter registered email and wrong password\nClick 'Login'",
                  "Email: <registered email>\nPassword: wrong-password",
                  "'Your email or password is incorrect!' is displayed", "Yes"])
    sheet.append(["TC_BAD", "Wrong expectation", "Click 'Signup / Login'", None,
                  "'Done' is displayed", "Yes"])
    suite = tmp_path / "cases.xlsx"
    workbook.save(suite)
    monkeypatch.chdir(tmp_path)
    outcome = CliRunner().invoke(
        app, ["run", str(suite), "--headless", "--config", str(config_file)]
    )
    assert outcome.exit_code == 1, outcome.output
    assert "Passed: 1  Failed: 1  Needs review: 0" in outcome.output
    assert "Expected text was never visible: Done" in outcome.output
    results_file = next((tmp_path / "data" / "runs").glob("run_*/results.json"))
    results = json.loads(results_file.read_text(encoding="utf-8"))
    assert [item["status"] for item in results["tests"]] == ["passed", "failed"]

    selected = CliRunner().invoke(
        app, ["run", str(suite), "--headless", "--config", str(config_file), "--case", "TC_OK"]
    )
    assert selected.exit_code == 0, selected.output
    assert "Passed: 1  Failed: 0  Needs review: 0" in selected.output
