"""Real-browser verification of first-run discovery and cache-only reuse."""

from __future__ import annotations

import json
import re
import threading
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

import httpx2
import openpyxl
import pytest
import yaml
from openai import RateLimitError
from typer.testing import CliRunner

from selfheal.cli import app
from selfheal.healing.cache import LocatorCache
from selfheal.healing.locator import LocatorHealer, verified_locator
from selfheal.llm.openai_client import LocatorResponseError, OpenAILLMClient, OpenAIQuotaError
from selfheal.onboarding import (
    _conversion_pages,
    _convert_cases,
    generated_profile_path,
    normalize_site_url,
    prepare_site,
    run_new_site,
    site_name,
)
from selfheal.suite_converter import (
    converted_suite_path,
    model_cases,
    recover_passed_conversions,
    remove_generated_execute_column,
    review_cases,
    save_converted_suite,
    validate_conversion,
)
from selfheal.suite_loader import read_excel_tests
from selfheal.telemetry.healing_log import HealingEventStore


@pytest.fixture
def local_site():
    state = {"button_id": "continue-v1", "title": "Done"}

    class Handler(BaseHTTPRequestHandler):
        def do_GET(self):
            if self.path == "/":
                body = (
                    "<html><body><input id='private-field' value='private-secret'>"
                    f"<a id='{state['button_id']}' href='/done'>Continue</a></body></html>"
                )
            elif self.path == "/done":
                body = f"<html><head><title>{state['title']}</title></head><body><h1>Done</h1></body></html>"
            else:
                self.send_error(404)
                return
            payload = body.encode("utf-8")
            self.send_response(200)
            self.send_header("Content-Type", "text/html; charset=utf-8")
            self.send_header("Content-Length", str(len(payload)))
            self.end_headers()
            self.wfile.write(payload)

        def log_message(self, *_args):
            pass

    server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        yield f"http://127.0.0.1:{server.server_port}", state
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)


@pytest.fixture
def form_site():
    state = {"email_id": "email-v1"}

    class Handler(BaseHTTPRequestHandler):
        def do_GET(self):
            if self.path == "/":
                body = '<html><body><a href="/login">Login</a></body></html>'
            elif self.path == "/login":
                body = f"""<html><body><section><h2>Login to your account</h2>
                    <form id="login-form">
                        <input id="{state['email_id']}" name="email" type="email"
                            placeholder="Email Address" required>
                        <input id="password" name="password" type="password" required>
                        <button id="login-button" type="submit">Login</button>
                    </form></section><section><h2>New User Signup!</h2>
                    <form id="signup-form"><input name="email" type="email"
                        placeholder="Email Address"><button type="submit">Signup</button>
                    </form></section>
                    <div id="login-status"></div>
                    <script>
                    document.querySelector('form').addEventListener('submit', event => {{
                        event.preventDefault();
                        const email = document.querySelector('[name=email]').value;
                        const password = document.querySelector('[name=password]').value;
                        if (email === 'demo@example.test' && password === 'demo-secret') {{
                            document.querySelector('#login-status').textContent = 'Signed in';
                        }}
                    }});
                    </script></body></html>"""
            else:
                self.send_error(404)
                return
            payload = body.encode("utf-8")
            self.send_response(200)
            self.send_header("Content-Type", "text/html; charset=utf-8")
            self.send_header("Content-Length", str(len(payload)))
            self.end_headers()
            self.wfile.write(payload)

        def log_message(self, *_args):
            pass

    server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        yield f"http://127.0.0.1:{server.server_port}", state
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)


def _workbook(path: Path, steps: str) -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["TestCase ID", "Title", "Test Step", "Test Data", "Expected Result"])
    sheet.append(["TC_001", "Continue to result", steps, "", ""])
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)


class FixtureModel:
    def __init__(self, state):
        self.state = state
        self.discovery_calls = 0
        self.repair_calls = 0

    def discover_site(self, pages, targets):
        self.discovery_calls += 1
        assert "private-secret" not in str(pages)
        assert pages[0]["page"] == "/"
        button_id = re.search(r'id="(continue-v\d+)"', pages[0]["html"]).group(1)
        locators = []
        for target in targets:
            if target == {"page": "/", "command": "CLICK", "target": "Continue"}:
                locators.append({"page": "/", "command": "CLICK", "target": "Continue",
                                 "selector": f"#{button_id}", "confidence": 0.99})
            else:
                assert target == {"page": "/done", "command": "ASSERT_VISIBLE", "target": "Done heading"}
                locators.append({"page": "/done", "command": "ASSERT_VISIBLE", "target": "Done heading",
                                 "selector": "h1", "confidence": 0.99})
        return {
            "site_name": "Local Fixture",
            "locators": locators,
        }

    def identify_element(self, page_tree, target, confidence_threshold):
        self.repair_calls += 1
        assert "private-secret" not in page_tree
        assert "Continue" in target
        return {"found": True, "locator": f"#{self.state['button_id']}", "confidence": 0.99}


def _result(run_dir: Path):
    return json.loads((run_dir / "results.json").read_text(encoding="utf-8"))


def test_direct_locator_suite_needs_no_discovery_or_api_key(local_site, tmp_path, monkeypatch):
    url, _state = local_site
    suite = tmp_path / "suites" / "direct.xlsx"
    _workbook(suite, "OPEN | /\nCLICK | css=#continue-v1\nASSERT_URL | /done")
    monkeypatch.setenv("OPENAI_API_KEY", "")
    monkeypatch.setenv("LANGSMITH_TRACING", "false")

    first = _result(run_new_site(url, suite, root=tmp_path))
    second = _result(run_new_site(url, suite, root=tmp_path))

    assert first["status"] == second["status"] == "passed"
    assert first["discovery"]["api_calls"] == second["discovery"]["api_calls"] == 0
    assert first["healing"]["ai_calls"] == second["healing"]["ai_calls"] == 0


def test_review_upload_rechecks_same_workbook_and_runs_converted_copy(
    local_site, tmp_path, monkeypatch, capsys,
):
    url, state = local_site
    suite = tmp_path / "data" / "uploads" / "fixture.xlsx"
    _workbook(suite, "OPEN | /\nCLICK | Continue\nASSERT_URL | /done")
    original_bytes = suite.read_bytes()

    class ReviewModel(FixtureModel):
        conversion_calls = 0

        def convert_site(self, pages, cases):
            self.conversion_calls += 1
            assert any("Continue" in page["html"] for page in pages)
            assert [case["id"] for case in cases] == ["TC_001"]
            return {"cases": [{
                "id": "TC_001", "steps": ["OPEN | /wrong"], "checks": [],
                "review_reason": "",
            }]}

    model = ReviewModel(state)
    monkeypatch.setenv("OPENAI_API_KEY", "offline-test-key")
    monkeypatch.setenv("LANGSMITH_TRACING", "false")

    first = _result(run_new_site(url, suite, root=tmp_path, client=model, review_upload=True))
    second = _result(run_new_site(url, suite, root=tmp_path, client=model, review_upload=True))
    output = capsys.readouterr().out
    converted = converted_suite_path(tmp_path, url, suite)

    assert first["status"] == second["status"] == "passed"
    assert first["conversion"]["api_calls"] == second["conversion"]["api_calls"] == 1
    assert first["discovery"]["api_calls"] == 1
    assert second["discovery"]["api_calls"] == 0
    assert second["tests"][0]["cache_hits"] == 1
    assert second["tests"][0]["cached_selectors"] == [{
        "page": "/", "command": "CLICK", "target": "Continue",
        "selector": "#continue-v1",
    }]
    assert "Cache used: 1 locator hit(s)" in output
    assert model.conversion_calls == 2
    assert Path(first["suite"]) == Path(second["suite"]) == converted
    assert suite.read_bytes() == original_bytes
    assert read_excel_tests(converted)[0]["step"] == "OPEN | /\nCLICK | Continue\nASSERT_URL | /done"


def test_review_snapshots_include_pages_named_by_open_commands(local_site):
    url, _state = local_site
    cases = [{"title": "Result", "step": "OPEN | /done\nASSERT_TITLE | Done"}]

    pages = _conversion_pages(url, cases)

    assert any(page["page"] == "/done" and "<h1" in page["html"] for page in pages)


def test_discovery_cache_survives_suite_switch(local_site, tmp_path, monkeypatch):
    url, state = local_site
    suite_a = tmp_path / "suites" / "a.xlsx"
    suite_b = tmp_path / "suites" / "b.xlsx"
    direct_suite = tmp_path / "suites" / "direct.xlsx"
    _workbook(suite_a, "OPEN | /\nCLICK | Continue\nASSERT_URL | /done")
    _workbook(suite_b, "OPEN | /\nCLICK | Continue\nASSERT_VISIBLE | Done heading")
    _workbook(direct_suite, "OPEN | /\nCLICK | css=#continue-v1\nASSERT_URL | /done")
    model = FixtureModel(state)
    monkeypatch.setenv("OPENAI_API_KEY", "offline-test-key")
    monkeypatch.setenv("LANGSMITH_TRACING", "false")

    first = _result(run_new_site(url, suite_a, root=tmp_path, client=model))
    monkeypatch.setenv("OPENAI_API_KEY", "")
    direct = _result(run_new_site(url, direct_suite, root=tmp_path))
    monkeypatch.setenv("OPENAI_API_KEY", "offline-test-key")
    second = _result(run_new_site(url, suite_b, root=tmp_path, client=model))
    third = _result(run_new_site(url, suite_a, root=tmp_path, client=model))
    cache = LocatorCache(tmp_path / "data" / "store" / "elements.db")
    app_name = site_name(url)

    assert [result["status"] for result in (first, direct, second, third)] == ["passed"] * 4
    assert [result["discovery"]["api_calls"] for result in (first, direct, second, third)] == [1, 0, 1, 0]
    assert model.discovery_calls == 2
    assert cache.get(app_name, "prod", "/", "generic-click:continue") == "#continue-v1"
    assert cache.get(app_name, "prod", "/done", "generic-assert_visible:done heading") == "h1"


def test_upload_suite_discovery_reuse_and_stale_repair(local_site, tmp_path, monkeypatch):
    url, state = local_site
    suite = tmp_path / "suites" / "fixture.xlsx"
    _workbook(suite, "OPEN | /\nCLICK | Continue\nASSERT_URL | /done")
    model = FixtureModel(state)
    monkeypatch.setenv("OPENAI_API_KEY", "offline-test-key")
    monkeypatch.setenv("LANGSMITH_TRACING", "false")
    monkeypatch.setattr(LocatorHealer, "_get_client", lambda _self, _key: model)

    first = _result(run_new_site(url, suite, root=tmp_path, client=model))
    profile_path = generated_profile_path(tmp_path, normalize_site_url(url))
    config = yaml.safe_load(profile_path.read_text(encoding="utf-8"))
    cache = LocatorCache(tmp_path / "data" / "store" / "elements.db")
    key = "generic-click:continue"
    app = site_name(url)
    assert first["status"] == "passed", first
    assert first["discovery"]["api_calls"] == 1
    assert first["healing"]["ai_calls"] == 0
    assert first["llm_calls"] == 1
    assert first["tests"][0]["steps"][1]["detail"].endswith("(cache)")
    assert config["locators_by_page"]["/"]["CLICK"]["Continue"] == "#continue-v1"
    assert config["llm"]["api_key_env"] == "OPENAI_API_KEY"
    assert "offline-test-key" not in profile_path.read_text(encoding="utf-8")
    assert cache.get(app, "prod", "/", key) == "#continue-v1"
    assert model.discovery_calls == 1 and model.repair_calls == 0

    second = _result(run_new_site(url, suite, root=tmp_path, client=model))
    assert second["status"] == "passed", second
    assert second["discovery"]["api_calls"] == 0
    assert second["healing"]["ai_calls"] == 0
    assert second["healing"]["cache_hits"] == 1
    assert model.discovery_calls == 1 and model.repair_calls == 0

    state["button_id"] = "continue-v2"
    third = _result(run_new_site(url, suite, root=tmp_path, client=model))
    assert third["status"] == "passed", third
    assert third["discovery"]["api_calls"] == 0
    assert third["healing"]["ai_calls"] == 1
    assert third["tests"][0]["steps"][1]["detail"].endswith("(healed)")
    assert cache.get(app, "prod", "/", key) == "#continue-v2"
    assert model.discovery_calls == 1 and model.repair_calls == 1

    fourth = _result(run_new_site(url, suite, root=tmp_path, client=model))
    assert fourth["status"] == "passed", fourth
    assert fourth["healing"]["ai_calls"] == 0
    assert fourth["healing"]["cache_hits"] == 1
    assert model.discovery_calls == 1 and model.repair_calls == 1

    expanded_suite = tmp_path / "suites" / "expanded.xlsx"
    _workbook(expanded_suite, "OPEN | /\nCLICK | Continue\nASSERT_VISIBLE | Done heading")
    expanded = _result(run_new_site(url, expanded_suite, root=tmp_path, client=model))
    assert expanded["status"] == "passed", expanded
    assert expanded["discovery"]["api_calls"] == 1
    assert model.discovery_calls == 2


def test_authorized_form_conversion_discovery_cache_reuse_and_healing(form_site, tmp_path, monkeypatch):
    url, state = form_site
    suite = tmp_path / "suites" / "form.xlsx"
    suite.parent.mkdir(parents=True)
    book = openpyxl.Workbook()
    book.active.append(["TestCase ID", "Title", "Test Step", "Test Data", "Expected Result"])
    book.active.append([
        "TC_001", "Sign in", "Open the login page, fill credentials and click Login",
        "Email: demo@example.test\nPassword: demo-secret", "'Signed in' is displayed",
    ])
    book.save(suite)

    class FormModel:
        def __init__(self):
            self.conversion_calls = 0
            self.discovery_calls = 0
            self.healing_calls = 0

        def convert_site(self, pages, cases):
            self.conversion_calls += 1
            assert "demo@example.test" not in str(cases)
            assert "demo-secret" not in str(cases)
            assert cases[0]["data_keys"] == ["email", "password"]
            return {"cases": [{
                "id": "TC_001",
                "steps": [
                    "OPEN | /login",
                    "FILL | Email Address in Login to your account | ${DATA:email}",
                    "FILL | Password in Login to your account | ${DATA:password}",
                    "CLICK | Login submit",
                    "ASSERT_VISIBLE | text=Signed in",
                ],
                "checks": [{"expected": "'Signed in' is displayed",
                            "assertion": "ASSERT_VISIBLE | text=Signed in"}],
            }]}

        def discover_site(self, pages, targets):
            self.discovery_calls += 1
            assert {page["page"] for page in pages} == {"/", "/login"}
            assert '&quot;form_id&quot;: &quot;login-form&quot;' in pages[1]["html"]
            assert "data-form-id" not in pages[1]["html"]
            assert {target["target"] for target in targets} == {
                "Email Address in Login to your account",
                "Password in Login to your account", "Login submit",
            }
            return {"site_name": "Form Fixture", "locators": [
                {"page": "/login", "command": "FILL",
                 "target": "Email Address in Login to your account",
                 "selector": f"#{state['email_id']}", "confidence": 0.99},
                {"page": "/login", "command": "FILL",
                 "target": "Password in Login to your account",
                 "selector": "#password", "confidence": 0.99},
                {"page": "/login", "command": "CLICK", "target": "Login submit",
                 "selector": "#login-button", "confidence": 0.99},
            ]}

        def identify_element(self, page_tree, target, confidence_threshold):
            self.healing_calls += 1
            assert "demo-secret" not in page_tree
            assert "Email Address" in target
            return {"found": True, "locator": f"#{state['email_id']}", "confidence": 0.99}

    model = FormModel()
    monkeypatch.setenv("OPENAI_API_KEY", "offline-test-key")
    monkeypatch.setenv("LANGSMITH_TRACING", "false")
    monkeypatch.setattr(LocatorHealer, "_get_client", lambda _self, _key: model)

    first = _result(run_new_site(
        url, suite, root=tmp_path, client=model, allow_site_actions=True,
    ))
    assert first["status"] == "passed", first
    assert first["conversion"]["api_calls"] == 1
    assert first["discovery"]["api_calls"] == 1
    assert first["healing"]["ai_calls"] == 0
    assert first["llm_calls"] == 2
    assert first["headless"] is True
    assert [step["status"] for step in first["tests"][0]["steps"]] == ["passed"] * 5
    assert first["suite"].startswith(str(tmp_path / "suites" / "converted"))

    second = _result(run_new_site(
        url, suite, root=tmp_path, client=model, allow_site_actions=True,
    ))
    assert second["status"] == "passed"
    assert second["llm_calls"] == 0
    assert second["healing"]["cache_hits"] == 3
    assert model.conversion_calls == 1 and model.discovery_calls == 1

    state["email_id"] = "email-v2"
    third = _result(run_new_site(
        url, suite, root=tmp_path, client=model, allow_site_actions=True,
    ))
    assert third["status"] == "passed", third
    assert third["healing"]["ai_calls"] == 1
    assert model.healing_calls == 1
    cache = LocatorCache(tmp_path / "data" / "store" / "elements.db")
    assert cache.get(
        site_name(url), "prod", "/login",
        "generic-fill:email address in login to your account",
    ) == "#email-v2"

    fourth = _result(run_new_site(
        url, suite, root=tmp_path, client=model, allow_site_actions=True,
    ))
    assert fourth["status"] == "passed"
    assert fourth["llm_calls"] == 0
    assert model.healing_calls == 1


def test_duplicate_form_fields_are_scoped_in_discovery_and_reused(form_site, tmp_path, monkeypatch):
    url, _state = form_site
    suite = tmp_path / "suites" / "duplicate-forms.xlsx"
    _workbook(suite, "\n".join([
        "OPEN | /login",
        "FILL | Email Address in Login to your account | demo@example.test",
        "FILL | Password in Login to your account | demo-secret",
        "CLICK | Login",
        "ASSERT_TEXT | css=#login-status | Signed in",
    ]))

    class Model:
        calls = 0

        def discover_site(self, _pages, targets):
            self.calls += 1
            selectors = {
                "Email Address in Login to your account": "input[name='email']",
                "Password in Login to your account": "input[name='password']",
                "Login": "button[type='submit']",
            }
            return {"site_name": "Two Forms", "locators": [
                {**target, "selector": selectors[target["target"]], "confidence": 0.99}
                for target in targets if target["target"] in selectors
            ]}

    model = Model()
    monkeypatch.setenv("OPENAI_API_KEY", "offline-test-key")
    monkeypatch.setenv("LANGSMITH_TRACING", "false")
    first = _result(run_new_site(url, suite, root=tmp_path, client=model))
    assert first["status"] == "passed", first
    assert first["healing"]["ai_calls"] == 0
    assert model.calls == 1
    cache = LocatorCache(tmp_path / "data" / "store" / "elements.db")
    selector = cache.get(
        site_name(url), "prod", "/login",
        "generic-fill:email address in login to your account",
    )
    assert selector == "form[id=\"login-form\"] input[name='email']"
    assert cache.get(site_name(url), "prod", "/login", "generic-click:login") == (
        "button[type='submit']:has-text(\"Login\")"
    )
    second = _result(run_new_site(url, suite, root=tmp_path, client=model))
    assert second["status"] == "passed", second
    assert second["llm_calls"] == 0
    assert model.calls == 1


def test_form_heading_scopes_prefix_and_suffix_targets(form_site):
    from playwright.sync_api import sync_playwright

    url, state = form_site
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=True, channel="chromium")
        try:
            page = browser.new_page()
            page.goto(url + "/login")
            page.evaluate(
                "document.body.insertAdjacentHTML('afterbegin', "
                "'<nav><a href=\"/login\">Signup / Login</a></nav>')"
            )
            login = verified_locator(
                page, "input[name='email']", "FILL",
                "Email Address in Login to your account", allow_scope=True,
            )
            assert login is not None
            assert login[1] == "form[id=\"login-form\"] input[name='email']"
            signup = verified_locator(
                page, "input[name='email']", "FILL",
                "New User Signup! Email Address", allow_scope=True,
            )
            assert signup is not None
            assert signup[1] == "form[id=\"signup-form\"] input[name='email']"
            assert verified_locator(
                page, f"#{state['email_id']}", "FILL", "New User Signup! Email Address",
                allow_scope=True,
            ) is None
            assert verified_locator(
                page, "nav a", "CLICK", "Login", allow_scope=True,
            ) is None
            assert verified_locator(page, "#login-button", "CLICK", "Login") is not None
        finally:
            browser.close()


def test_wrong_form_cache_is_healed_to_requested_form(form_site, tmp_path, monkeypatch):
    from playwright.sync_api import sync_playwright

    url, _state = form_site
    monkeypatch.setenv("OPENAI_API_KEY", "offline-test-key")
    cache = LocatorCache(tmp_path / "elements.db")
    cache_key = "generic-fill:login to your account email address"
    cache.put("fixture", "prod", "/login", cache_key, "form[id='signup-form'] input[name=email]")

    class Model:
        calls = 0

        def identify_element(self, _tree, _target, confidence_threshold=0.85):
            self.calls += 1
            return {"found": True, "locator": "input[name=email]", "confidence": 0.99}

    model = Model()
    config = {"app_name": "fixture", "resolver": {"ai_enabled": True},
              "healing": {"enabled": True, "auto_approve": True}}
    healer = LocatorHealer(
        config, "prod", langsmith=object(), cache=cache, client=model,
        event_store=HealingEventStore(tmp_path / "healing_events.db"),
    )
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=True, channel="chromium")
        try:
            page = browser.new_page()
            page.goto(url + "/login")
            match = healer.resolve(
                page, "Login to your account Email Address", kind="generic-fill",
            )
            assert match is not None and match.source == "healed"
            match.element.fill("demo@example.test")
            healer.remember_success(match)
            assert page.locator("#login-form input[name=email]").input_value() == "demo@example.test"
            assert page.locator("#signup-form input[name=email]").input_value() == ""
            assert cache.get("fixture", "prod", "/login", cache_key) == (
                'form[id="login-form"] input[name=email]'
            )
            assert model.calls == 1
        finally:
            browser.close()


def test_locator_validation_rejects_wrong_assertion_text_and_visible_target(local_site):
    from playwright.sync_api import sync_playwright

    url, state = local_site
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=True, channel="chromium")
        try:
            page = browser.new_page()
            page.goto(url)
            button = f"#{state['button_id']}"
            assert verified_locator(
                page, button, "ASSERT_TEXT", "result heading", expected_text="Done",
            ) is None
            assert verified_locator(page, button, "ASSERT_VISIBLE", "product list") is None
            page.goto(url + "/done")
            assert verified_locator(
                page, "h1", "ASSERT_TEXT", "Done heading", expected_text="Done",
            ) is not None
        finally:
            browser.close()


def test_exact_visible_text_resolves_without_ai_or_sensitive_cache(local_site, tmp_path, monkeypatch):
    from playwright.sync_api import sync_playwright

    url, _state = local_site
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    healer = LocatorHealer(
        {"app_name": "fixture", "resolver": {"ai_enabled": True},
         "healing": {"enabled": True, "auto_approve": True}},
        "prod", langsmith=object(),
        cache=LocatorCache(tmp_path / "elements.db"),
        event_store=HealingEventStore(tmp_path / "healing_events.db"),
    )
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=True, channel="chromium")
        try:
            page = browser.new_page()
            page.goto(url + "/done")
            match = healer.resolve(
                page, "Done heading", kind="generic-assert_text", expected_text="Done",
            )
            assert match is not None and match.source == "text"
            assert match.element.inner_text() == "Done"
            assert healer.ai_calls == 0
            assert healer.cache.get("fixture", "prod", "/done", "generic-assert_text:done heading") is None
        finally:
            browser.close()


def test_wait_visible_uses_unique_heading_text_without_ai(local_site, tmp_path, monkeypatch):
    from playwright.sync_api import sync_playwright

    from selfheal.executor.generic import GenericStepExecutor

    url, _state = local_site
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    config = {
        "app_name": "fixture", "environments": {"prod": {"base_url": url}},
        "resolver": {"ai_enabled": True, "discovery_mode": "batch_cached"},
        "healing": {"enabled": True, "auto_approve": True},
    }
    healer = LocatorHealer(
        config, "prod", langsmith=object(),
        cache=LocatorCache(tmp_path / "elements.db"),
        event_store=HealingEventStore(tmp_path / "healing_events.db"),
    )
    executor = GenericStepExecutor(config, "prod", tmp_path / "suite.xlsx", 5000, healer)
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=True, channel="chromium")
        try:
            page = browser.new_page()
            page.goto(url + "/done")
            assert executor.execute(page, "WAIT_VISIBLE | Done", {}) == "Waited for Done to become visible"
            assert healer.ai_calls == 0
            page.set_content("<h1>Done</h1><h2>Done</h2>")
            assert healer.resolve(page, "Done", kind="generic-wait_visible", allow_ai=False) is None
        finally:
            browser.close()


def test_converter_rejects_wrong_data_field_and_assertion_narrative():
    source = [{
        "id": "TC_001", "title": "Register", "step": "Fill First Name and submit",
        "data": "Name: Demo\nAddress: Main Street",
        "expected": "'Thank you for your review.' is displayed", "row": 2,
    }]
    payload = {"cases": [{
        "id": "TC_001",
        "steps": ["OPEN | /", "FILL | First Name | ${DATA:address}",
                  "ASSERT_TEXT | confirmation | Thank you for your review."],
        "checks": [{"expected": source[0]["expected"],
                    "assertion": "ASSERT_TEXT | confirmation | Thank you for your review."}],
    }]}
    result = validate_conversion(source, payload, "https://example.test", allow_site_actions=True)
    assert "wrong Test Data" in result[0]["conversion_note"]
    assert not result[0]["execute_partial"]

    payload["cases"][0]["steps"][1] = "FILL | First Name | ${DATA:name}"
    payload["cases"][0]["steps"][2] = (
        "ASSERT_TEXT | confirmation | Message 'Thank you for your review.' is displayed"
    )
    payload["cases"][0]["checks"][0]["assertion"] = payload["cases"][0]["steps"][2]
    result = validate_conversion(source, payload, "https://example.test", allow_site_actions=True)
    assert "quoted UI text" in result[0]["conversion_note"]


def test_converter_names_missing_keys_and_placeholder_values():
    source = [{
        "id": "TC_001", "title": "Login", "step": "Log in",
        "data": "Email: <registered email>\nPassword: secret",
        "expected": "'Done' title is displayed", "row": 2,
    }]
    proposal = {"cases": [{
        "id": "TC_001", "steps": [
            "OPEN | /login", "FILL | username | ${DATA:username}",
            "FILL | first name | ${DATA:first_name}", "ASSERT_TITLE | Done",
        ],
        "checks": [{"clause_index": 0, "assertion": "ASSERT_TITLE | Done"}],
    }]}
    converted = validate_conversion(source, proposal, "https://example.test", allow_site_actions=True)
    assert converted[0]["conversion_note"] == (
        "Generated command references an unknown Test Data field: first_name, username"
    )

    proposal["cases"][0]["steps"][1:3] = ["FILL | email | ${DATA:email}"]
    converted = validate_conversion(source, proposal, "https://example.test", allow_site_actions=True)
    assert converted[0]["conversion_note"] == "Generated command requires a real Test Data value: email"


def test_converter_sends_numbered_claims_without_data_values():
    source = [{"id": "TC_001", "title": "Search", "step": "Search for Top",
               "expected": "Heading is visible; results contain Top",
               "data": "Email: private@example.test", "row": 2}]
    case = model_cases(source)[0]
    assert case["expected_clauses"] == ["Heading is visible", "results contain Top"]
    assert case["data_keys"] == ["email"]
    assert "private@example.test" not in str(case)


def test_converter_accepts_indexed_assertion_mapping():
    source = [{"id": "TC_001", "title": "Done", "step": "Open Done",
               "expected": "'Done' title is displayed", "data": "", "row": 2}]
    proposal = {"cases": [{"id": "TC_001", "steps": ["OPEN | /done", "ASSERT_TITLE | Done"],
                           "checks": [{"clause_index": 0, "assertion": "ASSERT_TITLE | Done"}],
                           "review_reason": ""}]}
    converted = validate_conversion(source, proposal, "https://example.test")
    assert not converted[0]["conversion_note"]


def test_generated_upload_requires_real_absolute_file_path():
    source = [{"id": "TC_001", "title": "Contact", "step": "Upload file",
               "expected": "'Done' title is displayed",
               "data": "File: sample.txt (<1 MB)", "row": 2}]
    proposal = {"cases": [{"id": "TC_001", "steps": [
        "OPEN | /", "UPLOAD | file input | ${DATA:file}", "ASSERT_TITLE | Done",
    ], "checks": [{"clause_index": 0, "assertion": "ASSERT_TITLE | Done"}],
        "review_reason": ""}]}
    converted = validate_conversion(source, proposal, "https://example.test", allow_site_actions=True)
    assert "existing absolute file path" in converted[0]["conversion_note"]
    assert not converted[0]["execute_partial"]


def test_converter_rejects_browser_title_for_page_heading():
    source = [{"id": "TC_001", "title": "Category", "step": "Open category",
               "expected": "Page titled 'WOMEN - TOPS PRODUCTS' is shown",
               "data": "", "row": 2}]
    proposal = {"cases": [{"id": "TC_001", "steps": [
        "OPEN | /category", "ASSERT_TITLE | WOMEN - TOPS PRODUCTS",
    ], "checks": [{"clause_index": 0, "assertion": "ASSERT_TITLE | WOMEN - TOPS PRODUCTS"}],
        "review_reason": ""}]}
    converted = validate_conversion(source, proposal, "https://example.test")
    assert "does not verify" in converted[0]["conversion_note"]
    assert not converted[0]["execute_partial"]


def test_failed_assertion_does_not_trigger_locator_healing(local_site, tmp_path, monkeypatch):
    url, state = local_site
    suite = tmp_path / "suites" / "assertion.xlsx"
    _workbook(suite, "OPEN | /\nCLICK | Continue\nASSERT_TITLE | Expected title")
    model = FixtureModel(state)
    monkeypatch.setenv("OPENAI_API_KEY", "offline-test-key")
    monkeypatch.setenv("LANGSMITH_TRACING", "false")
    monkeypatch.setattr(LocatorHealer, "_get_client", lambda _self, _key: model)

    result = _result(run_new_site(url, suite, root=tmp_path, client=model))
    assert result["status"] == "failed"
    assert result["tests"][0]["steps"][-1]["status"] == "failed"
    assert model.discovery_calls == 1 and model.repair_calls == 0


def test_natural_suite_converts_discovers_reuses_and_heals(local_site, tmp_path, monkeypatch, capsys):
    url, state = local_site
    suite = tmp_path / "suites" / "prose.xlsx"
    book = openpyxl.Workbook()
    book.active.append(["TestCase ID", "Title", "Test Step", "Test Data", "Expected Result"])
    book.active.append([
        "TC_001", "Continue", "1. Open home page\n2. Click Continue",
        "Private: never-send-this-value", "'Done' title is displayed",
    ])
    suite.parent.mkdir(parents=True)
    book.save(suite)

    class ConversionModel(FixtureModel):
        conversion_calls = 0

        def convert_site(self, pages, cases):
            self.conversion_calls += 1
            assert "never-send-this-value" not in str(cases)
            assert "private-secret" not in str(pages)
            return {
                "site_name": "Local Fixture",
                "cases": [{
                    "id": "TC_001",
                    "steps": ["OPEN | /", "CLICK | Continue", "ASSERT_TITLE | Done"],
                    "checks": [{"expected": "'Done' title is displayed", "assertion": "ASSERT_TITLE | Done"}],
                }],
                "locators": [{"page": "/", "command": "CLICK", "target": "Continue",
                              "selector": f"#{state['button_id']}", "confidence": 0.99}],
            }

    model = ConversionModel(state)
    monkeypatch.setenv("OPENAI_API_KEY", "offline-test-key")
    monkeypatch.setenv("LANGSMITH_TRACING", "false")
    monkeypatch.setattr(LocatorHealer, "_get_client", lambda _self, _key: model)
    first = _result(run_new_site(url, suite, root=tmp_path, client=model))
    converted_path = converted_suite_path(tmp_path, url, suite)
    assert first["status"] == "passed", first
    assert first["discovery"]["api_calls"] == 1
    assert first["conversion"]["api_calls"] == 1
    assert first["llm_calls"] == 2
    assert first["total"] == 1 and model.conversion_calls == 1 and model.discovery_calls == 1
    assert "HEALED:" not in capsys.readouterr().out
    assert converted_path.parent.parent == tmp_path / "suites" / "converted"
    assert converted_path.is_file() and not review_cases(converted_path)
    book = openpyxl.load_workbook(converted_path, read_only=True)
    try:
        assert "Execute" not in [cell.value for cell in book["Framework Cases"][1]]
    finally:
        book.close()

    second = _result(run_new_site(url, suite, root=tmp_path, client=model))
    assert second["status"] == "passed", second
    assert second["discovery"]["api_calls"] == 0
    assert second["healing"]["ai_calls"] == 0
    assert model.conversion_calls == 1
    assert "HEALED:" not in capsys.readouterr().out

    state["button_id"] = "continue-v2"
    third = _result(run_new_site(url, suite, root=tmp_path, client=model))
    assert third["status"] == "passed", third
    assert third["healing"]["ai_calls"] == 1
    assert third["tests"][0]["healing_ai_calls"] == 1
    assert third["tests"][0]["healed_steps"] == ["CLICK | Continue"]
    third_output = capsys.readouterr().out
    assert "HEALED: CLICK | Continue" in third_output
    assert "Healing AI calls: 1" in third_output
    assert model.conversion_calls == 1 and model.repair_calls == 1


def test_retry_reviewed_cases_keeps_valid_conversion_and_repairs_only_blocked(local_site, tmp_path, monkeypatch):
    url, state = local_site
    suite = tmp_path / "suite.xlsx"
    book = openpyxl.Workbook()
    book.active.append(["TestCase ID", "Title", "Test Step", "Test Data", "Expected Result"])
    book.active.append(["TC_001", "Continue", "Click Continue", "", "'Done' title is displayed"])
    book.active.append(["TC_002", "Direct", "Open Done", "", "'Done' title is displayed"])
    book.save(suite)

    class RetryModel(FixtureModel):
        conversion_calls = 0

        def convert_site(self, _pages, cases):
            self.conversion_calls += 1
            proposals = []
            for case in cases:
                if case["id"] == "TC_001":
                    steps = ["OPEN | /", "CLICK | Continue", "ASSERT_TITLE | Done"]
                elif self.conversion_calls <= 2:
                    proposals.append({"id": "TC_002", "steps": [], "checks": [],
                                      "review_reason": "Needs a better assertion"})
                    continue
                else:
                    steps = ["OPEN | /done", "ASSERT_TITLE | Done"]
                proposals.append({"id": case["id"], "steps": steps,
                                  "checks": [{"clause_index": 0,
                                              "assertions": ["ASSERT_TITLE | Done"]}],
                                  "review_reason": ""})
            return {"cases": proposals}

    model = RetryModel(state)
    monkeypatch.setenv("OPENAI_API_KEY", "offline-test-key")
    monkeypatch.setenv("LANGSMITH_TRACING", "false")
    monkeypatch.setattr(LocatorHealer, "_get_client", lambda _self, _key: model)
    first = _result(run_new_site(url, suite, root=tmp_path, client=model))
    assert (first["passed"], first["needs_review"]) == (1, 1)
    path = converted_suite_path(tmp_path, url, suite)
    accepted_step = read_excel_tests(path)[0]["step"]
    assert model.conversion_calls == 2

    repeat = _result(run_new_site(url, suite, root=tmp_path, client=model))
    assert repeat["conversion"]["api_calls"] == 0
    assert model.conversion_calls == 2

    repaired = _result(run_new_site(
        url, suite, root=tmp_path, client=model, retry_reviewed_cases=True,
    ))
    assert (repaired["passed"], repaired["needs_review"]) == (2, 0)
    assert repaired["conversion"]["api_calls"] == 1
    assert model.conversion_calls == 3
    assert read_excel_tests(path)[0]["step"] == accepted_step


def test_discovery_follows_a_bare_css_link_before_resolving_form_fields(form_site, tmp_path, monkeypatch):
    url, state = form_site
    tests = [{"id": "TC_001", "title": "Login page", "data": "Email: demo@example.test",
              "expected": "", "step": "OPEN | /\nCLICK | a[href='/login']\n"
              "FILL | Email Address in Login to your account | ${DATA:email}\n"
              "ASSERT_URL | /login"}]

    class DiscoveryModel:
        def discover_site(self, _pages, targets):
            assert targets == [{"page": "/login", "command": "FILL",
                                "target": "Email Address in Login to your account"}]
            return {"site_name": "Form Fixture", "locators": [{
                "page": "/login", "command": "FILL",
                "target": "Email Address in Login to your account",
                "selector": f"#{state['email_id']}", "confidence": 0.99,
            }]}

    monkeypatch.setenv("OPENAI_API_KEY", "offline-test-key")
    monkeypatch.setenv("LANGSMITH_TRACING", "false")
    prepared = prepare_site(url, tests, root=tmp_path, client=DiscoveryModel())
    assert prepared.discovery_calls == 1
    cache = LocatorCache(tmp_path / "data/store/elements.db")
    assert cache.get(prepared.app_name, "prod", "/login",
                     "generic-fill:email address in login to your account") == f"#{state['email_id']}"


def test_legacy_generated_execute_column_migrates_to_metadata(tmp_path):
    path = tmp_path / "converted.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Framework Cases"
    sheet.append([
        "TestCase ID", "Title", "Test Step", "Test Data", "Expected Result",
        "Execute", "Conversion Note", "Original Test Step", "Original Expected Result",
    ])
    sheet.append(["TC_001", "Complete", "OPEN | /\nASSERT_URL | /", "", "", "Yes", "", "Open", "URL matches"])
    sheet.append([
        "TC_002", "Partial", "OPEN | /\nASSERT_URL | /", "", "", "Yes",
        "Timing is not verified", "Open", "Loads quickly",
    ])
    sheet.append([
        "TC_003", "Blocked", "OPEN | /", "", "", "No",
        "Missing assertion", "Open", "Visible",
    ])
    book.create_sheet("Original Cases").append([
        "TestCase ID", "Title", "Test Step", "Test Data", "Expected Result",
    ])
    meta = book.create_sheet("Framework Meta")
    meta.append(["Converter Version", "8"])
    meta.append(["Site Actions Authorized", False])
    book.save(path)

    assert remove_generated_execute_column(path)
    assert not remove_generated_execute_column(path)
    migrated = openpyxl.load_workbook(path, read_only=True)
    try:
        assert "Execute" not in [cell.value for cell in migrated["Framework Cases"][1]]
        assert ("TC_002", "partial") in list(migrated["Framework Meta"].values)
    finally:
        migrated.close()
    assert [case["id"] for case in read_excel_tests(path)] == ["TC_001", "TC_002"]
    assert [case["id"] for case in review_cases(path)] == ["TC_003"]


def test_reconversion_restores_only_validated_same_source_passes(tmp_path):
    source = tmp_path / "source.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["TestCase ID", "Title", "Test Step", "Test Data", "Expected Result"])
    sheet.append(["TC_001", "Continue", "Click Continue", "", "'Done' title is displayed"])
    sheet.append(["TC_002", "Timing", "Open home", "", "Home loads within acceptable time"])
    book.save(source)
    cases = read_excel_tests(source)
    bad = validate_conversion(cases, {"cases": [
        {"id": "TC_001", "steps": ["OPEN | /"], "checks": []},
        {"id": "TC_002", "steps": ["OPEN | /"], "checks": []},
    ]}, "https://example.test")
    converted = converted_suite_path(tmp_path, "https://example.test", source)
    save_converted_suite(source, converted, bad)
    result_file = tmp_path / "data" / "runs" / "run_20261010_000001" / "results.json"
    result_file.parent.mkdir(parents=True)
    result = {
        "suite": str(converted), "base_url": "https://wrong.test",
        "tests": [{
            "id": "TC_001", "status": "passed",
            "steps": [{"status": "passed", "text": step} for step in (
                "OPEN | /", "CLICK | Continue", "ASSERT_TITLE | Done",
            )],
        }, {
            "id": "TC_002", "status": "passed",
            "steps": [{"status": "passed", "text": "ASSERT_URL | /"}],
        }],
    }
    result_file.write_text(json.dumps(result), encoding="utf-8")
    assert recover_passed_conversions(converted, root=tmp_path, base_url="https://example.test") == 0

    result["base_url"] = "https://example.test"
    result_file.write_text(json.dumps(result), encoding="utf-8")
    assert recover_passed_conversions(converted, root=tmp_path, base_url="https://example.test") == 1
    assert recover_passed_conversions(converted, root=tmp_path, base_url="https://example.test") == 0
    assert [case["id"] for case in read_excel_tests(converted)] == ["TC_001"]
    assert [case["id"] for case in review_cases(converted)] == ["TC_002"]


def test_onboarding_recovers_passed_case_after_bad_reconversion(local_site, tmp_path, monkeypatch):
    url, state = local_site
    source = tmp_path / "source.xlsx"
    book = openpyxl.Workbook()
    book.active.append(["TestCase ID", "Title", "Test Step", "Test Data", "Expected Result"])
    book.active.append(["TC_001", "Continue", "Click Continue", "", "'Done' title is displayed"])
    book.save(source)
    converted = converted_suite_path(tmp_path, url, source)
    result_file = tmp_path / "data" / "runs" / "run_20261010_000001" / "results.json"
    result_file.parent.mkdir(parents=True)
    result_file.write_text(json.dumps({
        "suite": str(converted), "base_url": url,
        "tests": [{"id": "TC_001", "status": "passed", "steps": [
            {"status": "passed", "text": step} for step in (
                "OPEN | /", "CLICK | Continue", "ASSERT_TITLE | Done",
            )
        ]}],
    }), encoding="utf-8")

    class BadConversionModel(FixtureModel):
        conversion_calls = 0

        def convert_site(self, _pages, _cases):
            self.conversion_calls += 1
            return {"cases": [{"id": "TC_001", "steps": ["OPEN | /"], "checks": []}]}

    model = BadConversionModel(state)
    monkeypatch.setenv("OPENAI_API_KEY", "offline-test-key")
    monkeypatch.setenv("LANGSMITH_TRACING", "false")
    first = _result(run_new_site(url, source, root=tmp_path, client=model))
    assert first["status"] == "passed", first
    assert first["passed"] == 1 and first["needs_review"] == 0
    assert [case["id"] for case in read_excel_tests(converted)] == ["TC_001"]
    assert model.conversion_calls == 2
    second = _result(run_new_site(url, source, root=tmp_path, client=model))
    assert second["status"] == "passed", second
    assert second["conversion"]["api_calls"] == 0
    assert model.conversion_calls == 2


def test_uploaded_incomplete_converted_workbook_recovers_original_cases(local_site, tmp_path, monkeypatch):
    url, state = local_site
    suite = tmp_path / "uploaded-converted.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Framework Cases"
    sheet.append([
        "TestCase ID", "Title", "Test Step", "Test Data", "Expected Result",
        "Execute", "Conversion Note", "Original Test Step", "Original Expected Result",
    ])
    sheet.append([
        "TC_001", "Continue", "OPEN | /\nWAIT_VISIBLE | h1", "", "", "yes",
        "Expected Result clause has no executable assertion", "Click Continue",
        "'Done' title is displayed",
    ])
    original = book.create_sheet("Original Cases")
    original.append(["TestCase ID", "Title", "Test Step", "Test Data", "Expected Result"])
    original.append(["TC_001", "Continue", "Click Continue", "", "'Done' title is displayed"])
    book.save(suite)

    class Model(FixtureModel):
        calls = 0

        def convert_site(self, _pages, cases):
            self.calls += 1
            assert cases[0]["steps"] == "Click Continue"
            return {
                "site_name": "Local Fixture",
                "cases": [{"id": "TC_001",
                           "steps": ["OPEN | /", "CLICK | Continue", "ASSERT_TITLE | Done"],
                           "checks": [{"expected": "'Done' title is displayed",
                                       "assertion": "ASSERT_TITLE | Done"}]}],
                "locators": [{"page": "/", "command": "CLICK", "target": "Continue",
                              "selector": f"#{state['button_id']}", "confidence": 0.99}],
            }

    model = Model(state)
    monkeypatch.setenv("OPENAI_API_KEY", "offline-test-key")
    monkeypatch.setenv("LANGSMITH_TRACING", "false")
    result = _result(run_new_site(url, suite, root=tmp_path, client=model))
    assert result["status"] == "passed", result
    assert model.calls == 1
    assert result["discovery"]["api_calls"] == 1
    repaired = openpyxl.load_workbook(result["suite"], read_only=True)
    try:
        assert repaired["Original Cases"]["C2"].value == "Click Continue"
    finally:
        repaired.close()


def test_action_only_conversion_gets_targeted_retry_and_reuses_it(local_site, tmp_path, monkeypatch):
    url, state = local_site
    suite = tmp_path / "prose.xlsx"
    book = openpyxl.Workbook()
    book.active.append(["TestCase ID", "Title", "Test Step", "Test Data", "Expected Result"])
    book.active.append(["TC_001", "Continue", "Click Continue", "", "'Done' title is displayed"])
    book.save(suite)

    class Model(FixtureModel):
        calls = 0

        def convert_site(self, _pages, _cases):
            self.calls += 1
            steps = ["OPEN | /", "CLICK | Continue"]
            checks = []
            if self.calls == 2:
                assert _cases[0]["previous_steps"] == steps
                assert _cases[0]["previous_checks"] == []
                assert _cases[0]["validation_feedback"]
                steps.append("ASSERT_TITLE | Done")
                checks = [{"expected": "'Done' title is displayed", "assertion": "ASSERT_TITLE | Done"}]
            return {
                "site_name": "Local Fixture",
                "cases": [{"id": "TC_001", "steps": steps, "checks": checks}],
                "locators": [{"page": "/", "command": "CLICK", "target": "Continue",
                              "selector": f"#{state['button_id']}", "confidence": 0.99}],
            }

    model = Model(state)
    monkeypatch.setenv("OPENAI_API_KEY", "offline-test-key")
    monkeypatch.setenv("LANGSMITH_TRACING", "false")
    first = _result(run_new_site(url, suite, root=tmp_path, client=model))
    assert first["status"] == "passed", first
    assert first["conversion"]["api_calls"] == 2
    assert first["discovery"]["api_calls"] == 1
    assert model.calls == 2
    second = _result(run_new_site(url, suite, root=tmp_path, client=model))
    assert second["status"] == "passed", second
    assert second["discovery"]["api_calls"] == 0
    assert model.calls == 2


def test_all_unverifiable_cases_receive_review_results(local_site, tmp_path, monkeypatch):
    url, state = local_site
    suite = tmp_path / "prose.xlsx"
    book = openpyxl.Workbook()
    book.active.append(["TestCase ID", "Title", "Test Step", "Test Data", "Expected Result"])
    book.active.append(["TC_001", "Home", "Open home", "", "Home loads within acceptable time"])
    book.save(suite)

    class Model(FixtureModel):
        calls = 0

        def convert_site(self, _pages, _cases):
            self.calls += 1
            return {"site_name": "Local Fixture", "cases": [{"id": "TC_001",
                    "steps": ["OPEN | /"], "checks": []}], "locators": []}

    model = Model(state)
    monkeypatch.setenv("OPENAI_API_KEY", "offline-test-key")
    monkeypatch.setenv("LANGSMITH_TRACING", "false")
    first = _result(run_new_site(url, suite, root=tmp_path, client=model))
    assert first["status"] == "needs_review", first
    assert first["total"] == 1 and first["needs_review"] == 1 and first["passed"] == 0
    assert generated_profile_path(tmp_path, url).is_file()
    assert first["conversion"]["api_calls"] == 1
    assert first["discovery"]["api_calls"] == 0
    assert model.calls == 1 and model.discovery_calls == 0
    second = _result(run_new_site(url, suite, root=tmp_path, client=model))
    assert second["status"] == "needs_review"
    assert second["discovery"]["api_calls"] == 0
    assert model.calls == 1


def test_late_page_assertion_locator_is_discovered_and_cached(local_site, tmp_path, monkeypatch):
    url, state = local_site
    suite = tmp_path / "prose.xlsx"
    book = openpyxl.Workbook()
    book.active.append(["TestCase ID", "Title", "Test Step", "Test Data", "Expected Result"])
    book.active.append(["TC_001", "Continue", "Click Continue", "", "'Done' heading is displayed"])
    book.save(suite)

    class Model(FixtureModel):
        calls = 0

        def convert_site(self, _pages, _cases):
            self.calls += 1
            return {
                "site_name": "Local Fixture",
                "cases": [{"id": "TC_001", "steps": ["OPEN | /", "CLICK | Continue",
                           "ASSERT_VISIBLE | Done heading"],
                           "checks": [{"expected": "'Done' heading is displayed",
                                       "assertion": "ASSERT_VISIBLE | Done heading"}]}],
                "locators": [{"page": "/", "command": "CLICK", "target": "Continue",
                              "selector": f"#{state['button_id']}", "confidence": 0.99}],
            }

        def identify_element(self, page_tree, target, confidence_threshold):
            self.repair_calls += 1
            assert "Done heading" in target and "Done" in page_tree
            return {"found": True, "locator": "h1", "confidence": 0.99}

    model = Model(state)
    monkeypatch.setenv("OPENAI_API_KEY", "offline-test-key")
    monkeypatch.setenv("LANGSMITH_TRACING", "false")
    monkeypatch.setattr(LocatorHealer, "_get_client", lambda _self, _key: model)
    first = _result(run_new_site(url, suite, root=tmp_path, client=model))
    assert first["status"] == "passed", first
    assert model.repair_calls == 0
    cache = LocatorCache(tmp_path / "data" / "store" / "elements.db")
    assert cache.get(site_name(url), "prod", "/done", "generic-assert_visible:done heading") == "h1"
    second = _result(run_new_site(url, suite, root=tmp_path, client=model))
    assert second["status"] == "passed", second
    assert second["healing"]["ai_calls"] == 0
    assert model.repair_calls == 0


def test_complex_prose_claim_is_reviewed_not_passed(local_site, tmp_path, monkeypatch):
    url, state = local_site
    suite = tmp_path / "suites" / "mixed.xlsx"
    book = openpyxl.Workbook()
    book.active.append(["TestCase ID", "Title", "Test Step", "Test Data", "Expected Result"])
    book.active.append(["TC_001", "Continue", "Click Continue", "", "'Done' title is displayed"])
    book.active.append(["TC_002", "Timing", "Open the home page", "", "Home loads within acceptable time"])
    book.active.append(["TC_003", "Order", "Place order", "", "Order placed"])
    suite.parent.mkdir(parents=True)
    book.save(suite)

    class MixedModel(FixtureModel):
        def convert_site(self, _pages, _cases):
            return {
                "site_name": "Local Fixture",
                "cases": [
                    {"id": "TC_001", "steps": ["OPEN | /", "CLICK | Continue", "ASSERT_TITLE | Done"],
                     "checks": [{"expected": "'Done' title is displayed", "assertion": "ASSERT_TITLE | Done"}]},
                    {"id": "TC_002", "steps": ["OPEN | /", "ASSERT_URL | /"],
                     "checks": [{"expected": "Home loads within acceptable time", "assertion": "ASSERT_URL | /"}]},
                    {"id": "TC_003", "steps": [], "checks": [],
                     "review_reason": "Payment requires manual review"},
                ],
                "locators": [{"page": "/", "command": "CLICK", "target": "Continue",
                              "selector": f"#{state['button_id']}", "confidence": 0.99}],
            }

    model = MixedModel(state)
    monkeypatch.setenv("OPENAI_API_KEY", "offline-test-key")
    monkeypatch.setenv("LANGSMITH_TRACING", "false")
    monkeypatch.setattr(LocatorHealer, "_get_client", lambda _self, _key: model)
    result = _result(run_new_site(url, suite, root=tmp_path, client=model))
    assert result["status"] == "needs_review", result
    assert result["total"] == 3 and result["passed"] == 1 and result["needs_review"] == 2
    assert {item["id"] for item in result["tests"]} == {"TC_001", "TC_002", "TC_003"}
    timing = next(item for item in result["tests"] if item["id"] == "TC_002")
    assert timing["status"] == "needs_review"
    assert "timing" in timing["error"].lower()
    assert timing["steps"], "Safe partial checks should execute in the browser"

    converted_path = converted_suite_path(tmp_path, url, suite)
    repeated = _result(run_new_site(url, converted_path, root=tmp_path, client=model))
    assert repeated["total"] == 3 and repeated["needs_review"] == 2
    assert repeated["discovery"]["api_calls"] == 0

    monkeypatch.chdir(tmp_path)
    existing = set((tmp_path / "data" / "runs").glob("run_*/results.json"))
    manual = CliRunner().invoke(app, [
        "run", str(converted_path), "--headless", "--config",
        str(generated_profile_path(tmp_path, url)),
    ])
    assert manual.exit_code == 1, manual.output
    assert "Passed: 1  Failed: 0  Needs review: 2" in manual.output
    manual_file, = set((tmp_path / "data" / "runs").glob("run_*/results.json")) - existing
    manual_result = _result(manual_file.parent)
    assert manual_result["total"] == 3
    assert {item["id"] for item in manual_result["tests"]} == {"TC_001", "TC_002", "TC_003"}
    assert manual_result["llm_calls"] == 0

    review_only = CliRunner().invoke(app, [
        "run", str(converted_path), "--headless", "--config",
        str(generated_profile_path(tmp_path, url)), "--case", "TC_003",
    ])
    assert review_only.exit_code == 1, review_only.output
    assert "Passed: 0  Failed: 0  Needs review: 1" in review_only.output


def test_converter_quarantines_sensitive_data_actions():
    source = [{
        "id": "TC_001", "title": "Login", "step": "Log in with the test password",
        "data": "Password: demo-secret", "expected": "'Logged in' title is displayed", "row": 2,
    }]
    proposal = {
        "cases": [{
            "id": "TC_001",
            "steps": ["OPEN | /login", "FILL | css=#password | ${DATA:Password}",
                      "CLICK | css=#submit", "ASSERT_TITLE | Logged in"],
            "checks": [{"expected": "'Logged in' title is displayed",
                        "assertion": "ASSERT_TITLE | Logged in"}],
        }],
    }
    converted = validate_conversion(source, proposal, "https://example.test")
    assert not converted[0]["execute_partial"]
    assert "Sensitive Test Data" in converted[0]["conversion_note"]


def test_converter_rejects_untrusted_environment_reference():
    source = [{"id": "TC_001", "title": "Submit", "step": "Submit the form",
               "data": "", "expected": "'Done' is displayed", "row": 2}]
    proposal = {"cases": [{
        "id": "TC_001",
        "steps": ["OPEN | /", "FILL | css=#message | ${ENV:OPENAI_API_KEY}",
                  "CLICK | css=#submit", "ASSERT_VISIBLE | text=Done"],
        "checks": [{"expected": "'Done' is displayed", "assertion": "ASSERT_VISIBLE | text=Done"}],
    }]}
    converted = validate_conversion(source, proposal, "https://example.test")
    assert not converted[0]["execute_partial"]
    assert "environment secrets" in converted[0]["conversion_note"]


def test_onboarding_rejects_invalid_conversion_without_profile(local_site, tmp_path, monkeypatch):
    url, state = local_site
    suite = tmp_path / "suites" / "unsupported.xlsx"
    _workbook(suite, "Click Continue then verify the next page")
    class InvalidModel(FixtureModel):
        def convert_site(self, _pages, _cases):
            return {"cases": []}

    model = InvalidModel(state)
    monkeypatch.setenv("OPENAI_API_KEY", "offline-test-key")
    with pytest.raises(ValueError, match="omitted test cases"):
        run_new_site(url, suite, root=tmp_path, client=model)
    assert model.discovery_calls == 0
    assert not generated_profile_path(tmp_path, url).exists()


def test_invalid_discovery_does_not_create_profile_or_cache(local_site, tmp_path, monkeypatch):
    url, state = local_site
    suite = tmp_path / "suites" / "invalid.xlsx"
    _workbook(suite, "OPEN | /\nCLICK | Continue\nASSERT_URL | /done")
    model = FixtureModel(state)
    model.discover_site = lambda _pages, _targets: {"site_name": "fixture", "locators": []}
    monkeypatch.setenv("OPENAI_API_KEY", "offline-test-key")
    monkeypatch.setenv("LANGSMITH_TRACING", "false")

    with pytest.raises(ValueError, match="could not verify"):
        run_new_site(url, suite, root=tmp_path, client=model)
    assert not generated_profile_path(tmp_path, url).exists()
    assert not (tmp_path / "data" / "store" / "elements.db").exists()


def test_discovery_client_makes_one_request_without_sdk_retries(monkeypatch):
    calls = []

    class StubCompletions:
        def create(self, **kwargs):
            calls.append(kwargs)
            return type("Response", (), {
                "choices": [type("Choice", (), {
                    "message": type("Message", (), {"content": "not-json"})(),
                })()],
            })()

    class StubOpenAI:
        def __init__(self, **kwargs):
            assert kwargs["api_key"] == "offline-test-key"
            assert kwargs["max_retries"] == 0
            self.chat = type("Chat", (), {"completions": StubCompletions()})()

    class NoTrace:
        def wrap_openai(self, client):
            return client

    monkeypatch.setattr("selfheal.llm.openai_client.OpenAI", StubOpenAI)
    client = OpenAILLMClient("offline-test-key", langsmith=NoTrace())
    with pytest.raises(LocatorResponseError, match="valid JSON"):
        client.discover_site([{"page": "/", "html": "<a>Continue</a>"}], [])
    assert len(calls) == 1


def test_conversion_client_uses_one_json_request(monkeypatch):
    calls = []

    class StubCompletions:
        def create(self, **kwargs):
            calls.append(kwargs)
            return type("Response", (), {
                "choices": [type("Choice", (), {
                    "message": type("Message", (), {
                        "content": '{"site_name":"Fixture","cases":[],"locators":[]}',
                    })(),
                })()],
            })()

    class StubOpenAI:
        def __init__(self, **kwargs):
            assert kwargs["api_key"] == "offline-test-key"
            assert kwargs["max_retries"] == 0
            self.chat = type("Chat", (), {"completions": StubCompletions()})()

    class NoTrace:
        def wrap_openai(self, client):
            return client

    monkeypatch.setattr("selfheal.llm.openai_client.OpenAI", StubOpenAI)
    client = OpenAILLMClient("offline-test-key", model="gpt-4o-mini", langsmith=NoTrace())
    result = client.convert_site([{"page": "/", "html": "<a>Continue</a>"}], [])
    assert result["site_name"] == "Fixture"
    assert len(calls) == 1
    assert calls[0]["response_format"]["type"] == "json_schema"
    assert calls[0]["response_format"]["json_schema"]["strict"] is True


@pytest.mark.parametrize("method", ["convert_site", "discover_site", "identify_element"])
def test_exhausted_credits_stop_openai_calls_without_retry(monkeypatch, method):
    calls = []
    request = httpx2.Request("POST", "https://api.openai.com/v1/chat/completions")
    response = httpx2.Response(429, request=request, json={
        "error": {"code": "credit_balance_exhausted", "type": "insufficient_quota"},
    })
    error = RateLimitError("No credits remaining", response=response, body=response.json()["error"])

    class StubCompletions:
        def create(self, **_kwargs):
            calls.append(1)
            raise error

    class StubOpenAI:
        def __init__(self, **_kwargs):
            self.chat = type("Chat", (), {"completions": StubCompletions()})()

    class NoTrace:
        def wrap_openai(self, client):
            return client

    monkeypatch.setattr("selfheal.llm.openai_client.OpenAI", StubOpenAI)
    client = OpenAILLMClient("offline-test-key", langsmith=NoTrace(), max_retries=3)
    with pytest.raises(OpenAIQuotaError, match="credits are exhausted"):
        if method == "convert_site":
            client.convert_site([], [])
        elif method == "discover_site":
            client.discover_site([], [])
        else:
            client.identify_element("[]", "Continue")
    assert len(calls) == 1


def test_conversion_retry_does_not_hide_exhausted_credits():
    class Model:
        calls = 0

        def convert_site(self, _pages, _cases):
            self.calls += 1
            if self.calls == 2:
                raise OpenAIQuotaError("OpenAI API credits are exhausted")
            return {"cases": [{"id": "TC_001", "steps": ["OPEN | /"], "checks": []}],
                    "locators": []}

    model = Model()
    source = [{"id": "TC_001", "title": "Home", "step": "Open home page",
               "expected": "Home is visible", "data": "", "row": 2}]
    with pytest.raises(OpenAIQuotaError, match="credits are exhausted"):
        _convert_cases(model, [], source, "https://example.test")
    assert model.calls == 2


def test_malformed_focused_retry_preserves_review_instead_of_running_actions():
    class Model:
        calls = 0

        def convert_site(self, _pages, _cases):
            self.calls += 1
            if self.calls == 2:
                return {}
            return {"cases": [{"id": "TC_001", "steps": ["OPEN | /", "CLICK | Continue"],
                               "checks": [], "review_reason": ""}]}

    source = [{"id": "TC_001", "title": "Continue", "step": "Click Continue",
               "expected": "Done page is displayed", "data": "", "row": 2}]
    model = Model()
    converted, calls = _convert_cases(model, [], source, "https://example.test")
    assert calls == model.calls == 2
    assert converted[0]["conversion_note"]
    assert not converted[0]["execute_partial"]


def test_onboard_cli_shows_quota_message_without_traceback(monkeypatch):
    def fail_quota(*_args, **_kwargs):
        raise OpenAIQuotaError("OpenAI API credits are exhausted")

    monkeypatch.setattr("selfheal.onboarding.run_new_site", fail_quota)
    result = CliRunner().invoke(app, ["onboard", "https://example.test", "tests.xlsx"])
    assert result.exit_code == 1
    assert "OpenAI API credits are exhausted" in result.output
    assert "Traceback" not in result.output


def test_onboard_cli_hides_discovery_call_count(monkeypatch, tmp_path):
    run_dir = tmp_path / "run"
    run_dir.mkdir()
    results = {
        "suite": "tests.xlsx",
        "discovery": {"profile": "generated.yaml", "api_calls": 1},
        "conversion": {"api_calls": 1},
        "healing": {"ai_calls": 0, "cache_hits": 2},
        "passed": 1, "failed": 0, "needs_review": 0, "status": "passed",
    }
    (run_dir / "results.json").write_text(json.dumps(results), encoding="utf-8")
    monkeypatch.setattr("selfheal.onboarding.run_new_site", lambda *_args, **_kwargs: run_dir)

    result = CliRunner().invoke(app, ["onboard", "https://example.test", "tests.xlsx"])

    assert result.exit_code == 0
    assert "Conversion API calls: 1" in result.output
    assert "Discovery API calls" not in result.output
    assert "Healing AI calls: 0" in result.output
    assert json.loads((run_dir / "results.json").read_text(encoding="utf-8"))["discovery"]["api_calls"] == 1


def test_manual_run_requires_explicit_config():
    result = CliRunner().invoke(app, ["run", "missing.xlsx"])
    assert result.exit_code == 1
    assert "pass --config" in result.output
