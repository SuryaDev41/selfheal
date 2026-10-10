"""Excel-backed history for suites launched from the Streamlit dashboard."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from selfheal.time_utils import IST, now_ist

DEFAULT_HISTORY_PATH = Path("data/dashboard/dashboard_history.xlsx")
SHEET_NAME = "Dashboard Runs"
HEADERS = (
    "Launch ID",
    "Started At (IST)",
    "Finished At (IST)",
    "Status",
    "Exit Code",
    "Website",
    "Suite",
    "Environment",
    "Headless",
    "Cases",
    "Console Output",
)
COLUMN_WIDTHS = (28, 22, 22, 16, 12, 24, 34, 14, 12, 34, 100)
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _naive_ist(value: datetime | None) -> datetime | None:
    return value.astimezone(IST).replace(tzinfo=None) if value else None


def _apply_layout(sheet) -> None:
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    for index, (header, width) in enumerate(zip(HEADERS, COLUMN_WIDTHS, strict=True), 1):
        cell = sheet.cell(row=1, column=index, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[1].height = 24


def _open_history(path: Path) -> tuple[Workbook, Any]:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(path) if path.is_file() else Workbook()
    sheet = workbook[SHEET_NAME] if SHEET_NAME in workbook.sheetnames else workbook.active
    sheet.title = SHEET_NAME
    _apply_layout(sheet)
    return workbook, sheet


def _finish_layout(sheet) -> None:
    last_column = get_column_letter(len(HEADERS))
    sheet.auto_filter.ref = f"A1:{last_column}{max(1, sheet.max_row)}"
    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column == len(HEADERS))
    for column in (2, 3):
        for cell in sheet.iter_cols(min_col=column, max_col=column, min_row=2, max_row=sheet.max_row):
            cell[0].number_format = "yyyy-mm-dd hh:mm:ss"


def start_dashboard_run(record: dict[str, Any], path: str | Path = DEFAULT_HISTORY_PATH) -> Path:
    """Append a pending dashboard launch and return its workbook path."""
    history_path = Path(path)
    workbook, sheet = _open_history(history_path)
    launch_id = record["launch_id"]
    known_ids = {sheet.cell(row=row, column=1).value for row in range(2, sheet.max_row + 1)}
    if launch_id not in known_ids:
        sheet.append(
            (
                launch_id,
                _naive_ist(record.get("started_at") or now_ist()),
                None,
                "running",
                None,
                record.get("website"),
                record.get("suite"),
                record.get("environment"),
                record.get("headless"),
                ", ".join(record.get("cases") or []) or "All enabled cases",
                "",
            )
        )
    _finish_layout(sheet)
    workbook.save(history_path)
    return history_path


def finish_dashboard_run(
    launch_id: str,
    *,
    exit_code: int,
    output: str,
    path: str | Path = DEFAULT_HISTORY_PATH,
) -> Path:
    """Mark a previously launched dashboard run as completed and save its output."""
    history_path = Path(path)
    workbook, sheet = _open_history(history_path)
    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row=row, column=1).value == launch_id:
            sheet.cell(row=row, column=3, value=_naive_ist(now_ist()))
            sheet.cell(row=row, column=4, value="passed" if exit_code == 0 else "failed")
            sheet.cell(row=row, column=5, value=exit_code)
            sheet.cell(row=row, column=11, value=output[-30000:])
            break
    _finish_layout(sheet)
    workbook.save(history_path)
    return history_path


def migrate_legacy_logs(directory: str | Path) -> Path:
    """Import old dashboard text logs once, then remove the migrated text files."""
    dashboard_dir = Path(directory)
    history_path = dashboard_dir / DEFAULT_HISTORY_PATH.name
    for log_path in dashboard_dir.glob("launch_*.log"):
        launch_id = log_path.stem
        try:
            started_at = datetime.strptime(launch_id.removeprefix("launch_"), "%Y%m%d_%H%M%S").replace(
                tzinfo=IST
            )
            output = log_path.read_text(encoding="utf-8", errors="replace")
            start_dashboard_run(
                {
                    "launch_id": launch_id,
                    "started_at": started_at,
                    "website": "Legacy dashboard launch",
                    "suite": "",
                    "environment": "",
                    "headless": "",
                    "cases": [],
                },
                history_path,
            )
            finish_dashboard_run(launch_id, exit_code=0, output=output, path=history_path)
            log_path.unlink()
        except (OSError, ValueError):
            continue
    return history_path
