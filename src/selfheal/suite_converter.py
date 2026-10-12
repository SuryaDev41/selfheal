"""Conservative conversion of uploaded prose cases into executable commands."""

from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path
from typing import Any
from uuid import uuid4

from openpyxl import Workbook, load_workbook

from selfheal.executor.browser import parse_test_data, split_steps
from selfheal.executor.generic import GenericStepExecutor

CONVERTER_VERSION = "16"
PRIVATE_VALUE = re.compile(r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}|\b\d{12,19}\b")
PRIVATE_ASSIGNMENT = re.compile(r"\b(password|token|secret|api[ _-]?key)\s*[:=]\s*\S+", re.I)
SENSITIVE_ACTION = re.compile(r"\bdelete\b|\bpay\b|place order|\bpurchase\b|\balert\b", re.I)
SENSITIVE_DATA_KEY = re.compile(r"password|card|cvc|cvv|token|secret|api.?key", re.I)
TWO_ARGUMENTS = {
    "FILL", "SELECT", "PRESS", "UPLOAD", "ASSERT_TEXT", "ASSERT_VALUE", "ASSERT_VALIDATION",
    "ASSERT_COUNT", "ASSERT_EACH_TEXT_CONTAINS", "ASSERT_COUNT_MATCHES",
    "ASSERT_NUMBER_EQUALS", "ASSERT_DATA_FIELDS",
}
THREE_ARGUMENTS = {"ASSERT_MULTIPLY"}
COLLECTION_ASSERTIONS = {
    "ASSERT_COUNT", "ASSERT_EACH_TEXT_CONTAINS", "ASSERT_COUNT_MATCHES",
    "ASSERT_NUMBER_EQUALS", "ASSERT_MULTIPLY", "ASSERT_DATA_FIELDS",
}


def needs_conversion(tests: list[dict[str, Any]]) -> bool:
    return any(
        str(case.get("expected") or "").strip()
        or any(not GenericStepExecutor.is_generic_step(step) for step in split_steps(case.get("step")))
        for case in tests
    )


def is_converted_suite(path: Path) -> bool:
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        return "Original Cases" in book.sheetnames
    finally:
        book.close()


def incomplete_converted_suite(
    path: Path, tests: list[dict[str, Any]], *, allow_site_actions: bool = False,
) -> bool:
    """Reconvert old generated workbooks or rows lacking an executable assertion."""
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        if "Original Cases" not in book.sheetnames:
            return False
        if "Framework Meta" not in book.sheetnames:
            return True
        meta = book["Framework Meta"]
        if meta["B1"].value != CONVERTER_VERSION or bool(meta["B2"].value) != allow_site_actions:
            return True
        return any(
            not any(step.strip().upper().startswith("ASSERT_") for step in split_steps(case["step"]))
            for case in tests
        )
    finally:
        book.close()


def converted_suite_path(
    root: Path, url: str, source: Path, *, allow_site_actions: bool = False,
) -> Path:
    from selfheal.onboarding import site_name

    mode = b"authorized" if allow_site_actions else b"restricted"
    digest = hashlib.sha256(CONVERTER_VERSION.encode() + mode + source.read_bytes()).hexdigest()[:16]
    return root / "suites" / "converted" / site_name(url) / f"{source.stem}-{digest}.xlsx"


def model_cases(tests: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Do not send spreadsheet data values or credentials to the model."""
    def redact(value: str) -> str:
        return PRIVATE_ASSIGNMENT.sub(r"\1: <redacted>", PRIVATE_VALUE.sub("<redacted>", value))

    output = []
    for case in tests:
        item = {
            "id": case["id"],
            "title": case["title"],
            "steps": redact(case["step"]),
            "expected": redact(case["expected"]),
            "expected_clauses": [redact(claim) for claim in _claims(case["expected"])],
            "data_keys": list(parse_test_data(case.get("data"))),
        }
        if case.get("validation_feedback"):
            item["validation_feedback"] = redact(str(case["validation_feedback"]))
        if case.get("previous_steps"):
            item["previous_steps"] = [redact(str(step)) for step in case["previous_steps"]]
        output.append(item)
    return output


def reviewed_attempts(path: Path) -> dict[str, dict[str, Any]]:
    """Return stored failed commands and feedback for an explicit repair retry."""
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book["Framework Cases"]
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows)
        attempts = {}
        for row in rows:
            values = dict(zip(headers, row, strict=False))
            note = str(values.get("Conversion Note") or "").strip()
            if not note:
                continue
            attempts[str(values["TestCase ID"])] = {
                "validation_feedback": note,
                "previous_steps": split_steps(values.get("Test Step")),
            }
        return attempts
    finally:
        book.close()


def _claims(expected: str) -> list[str]:
    return [claim.strip() for claim in expected.split(";") if claim.strip()]


def missing_time_limit(expected: str) -> bool:
    return bool(re.search(r"within acceptable time|\bload(?:s|ed)?\s+(?:in|within|under)\b", expected, re.I)) and not bool(
        re.search(r"\b\d+\s*(?:ms|milliseconds?|seconds?|s)\b", expected, re.I)
    )


def _data_key_issues(steps: list[str], raw_data: object) -> tuple[list[str], list[str]]:
    test_data = parse_test_data(raw_data)
    referenced = {
        key.strip().casefold()
        for step in steps
        for key in re.findall(r"\$\{DATA:([^}]+)}", step, re.I)
    }
    unknown = sorted(referenced - test_data.keys())
    placeholders = []
    for key in referenced & test_data.keys():
        value = re.sub(r"<timestamp>", "20000101000000000000", test_data[key], flags=re.I)
        if not value or re.search(r"<[^>]+>", value) or (
            "email" in key and not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", value)
        ):
            placeholders.append(key)
    return unknown, sorted(placeholders)


def _claim_issues(claim: str, assertions: list[str], steps: list[str]) -> list[str]:
    issues: list[str] = []
    commands = {assertion.split("|", 1)[0].strip().upper() for assertion in assertions}
    lower = claim.casefold()
    if re.search(r"within acceptable time|\bload(?:s|ed)?\s+(?:in|within|under)\b", lower):
        limit = re.search(r"\b(\d+)\s*(ms|milliseconds?|seconds?|s)\b", lower)
        if not limit:
            issues.append("Timing requires an explicit numeric time limit in Expected Result")
        else:
            milliseconds = int(limit.group(1)) * (1000 if limit.group(2).startswith("s") else 1)
            if f"ASSERT_LOAD_TIME | {milliseconds}" not in assertions:
                issues.append("Timing claim needs ASSERT_LOAD_TIME with the stated limit")
    if "validation" in lower or "html5" in lower:
        if "ASSERT_VALIDATION" not in commands:
            issues.append("Validation claim needs ASSERT_VALIDATION")
    if "count matches" in lower and "ASSERT_COUNT_MATCHES" not in commands:
        issues.append("Product count comparison needs ASSERT_COUNT_MATCHES")
    if "no products listed" in lower and not any(
        assertion.startswith("ASSERT_COUNT | ") and assertion.rsplit("|", 1)[-1].strip() == "0"
        for assertion in assertions
    ):
        issues.append("Empty product list needs ASSERT_COUNT with zero")
    if re.search(r"\bevery\b.*\bproducts?\b|\bproducts?\b.*\bonly\b|\bonly\b.*\bproducts?\b", lower) and "ASSERT_EACH_TEXT_CONTAINS" not in commands:
        issues.append("Each matching product needs ASSERT_EACH_TEXT_CONTAINS")
    cart_total = "total" in lower and any(word in lower for word in ("cart", "quantity", "rs.", "price"))
    if cart_total and "ASSERT_MULTIPLY" not in commands:
        issues.append("Cart total needs ASSERT_MULTIPLY")
    total = re.search(r"\btotal\s*=\s*(?:rs\.?\s*)?([\d,]+)", lower)
    if total and not any(
        assertion.startswith("ASSERT_NUMBER_EQUALS | ")
        and assertion.rsplit("|", 1)[-1].strip().replace(",", "") == total.group(1).replace(",", "")
        for assertion in assertions
    ):
        issues.append("Stated cart total needs ASSERT_NUMBER_EQUALS")
    price = re.search(r"\bprice\s*\(\s*rs\.?\s*([\d,]+)\s*\)", lower)
    if price and not any(
        assertion.startswith("ASSERT_NUMBER_EQUALS | ")
        and assertion.rsplit("|", 1)[-1].strip().replace(",", "") == price.group(1).replace(",", "")
        for assertion in assertions
    ):
        issues.append("Stated product price needs ASSERT_NUMBER_EQUALS")
    quantity = re.search(r"\bquantity\s+(\d+)\b", lower)
    if quantity and not any(
        assertion.startswith(("ASSERT_NUMBER_EQUALS | ", "ASSERT_VALUE | "))
        and assertion.rsplit("|", 1)[-1].strip() == quantity.group(1)
        for assertion in assertions
    ):
        issues.append("Stated quantity needs a numeric or input-value assertion")
    if "both products" in lower and len({
        assertion for assertion in assertions
        if assertion.startswith(("ASSERT_TEXT | ", "ASSERT_VISIBLE | "))
    }) < 2:
        issues.append("Both products need separate identity assertions")
    if "address matches" in lower and "ASSERT_DATA_FIELDS" not in commands:
        issues.append("Address comparison needs ASSERT_DATA_FIELDS")
    if "modal" in lower and "ASSERT_DIALOG" in commands:
        issues.append("An in-page modal needs ASSERT_TEXT or ASSERT_VISIBLE, not ASSERT_DIALOG")
    if re.search(r"\b(?:alert|dialog)\b", lower) and (
        "ASSERT_DIALOG" not in commands or not any(step.startswith("WATCH_DIALOG | ") for step in steps)
    ):
        issues.append("Browser dialog needs WATCH_DIALOG and ASSERT_DIALOG")
    if "after each" in lower:
        additions = [index for index, step in enumerate(steps) if step.startswith("CLICK | ") and "add" in step.casefold()]
        if len(additions) < 2 or any(
            not any(steps[index].startswith("ASSERT_") for index in range(start + 1, end))
            for start, end in zip(additions, [*additions[1:], len(steps)], strict=True)
        ):
            issues.append("Each add action needs its own following assertion")
    if re.search(r"\bheading\b|\bpage titled\b", lower):
        if not commands & {"ASSERT_TEXT", "ASSERT_VISIBLE"}:
            issues.append("Page heading needs an in-page text or visibility assertion")
    elif "title" in lower and any(word in lower for word in ("displayed", "shown", "contains")):
        if not commands & {"ASSERT_TITLE", "ASSERT_TEXT", "ASSERT_VISIBLE"}:
            issues.append("Title claim needs a matching assertion")
    native_validation = "validation" in lower or "html5" in lower
    if not native_validation and "ASSERT_VALIDATION" in commands and not commands & {
        "ASSERT_TEXT", "ASSERT_VISIBLE", "ASSERT_TITLE",
    }:
        issues.append("Server-rendered error needs an in-page text or visibility assertion")
    if not (native_validation and "ASSERT_VALIDATION" in commands) and any(
        word in lower for word in ("visible", "displayed", "shown", "appears")
    ):
        if not commands & {"ASSERT_VISIBLE", "ASSERT_TEXT", "ASSERT_TITLE", "ASSERT_DIALOG"}:
            issues.append("Visible outcome needs a visible text or element assertion")
    if "redirect" in lower or "url" in lower:
        if not commands & {"ASSERT_URL", "ASSERT_URL_CONTAINS"}:
            issues.append("Redirect needs ASSERT_URL or ASSERT_URL_CONTAINS")
    if "checkout page is not opened" in lower and "ASSERT_URL_NOT_CONTAINS" not in commands:
        issues.append("Blocked checkout needs ASSERT_URL_NOT_CONTAINS")
    return issues


def _claim_issue(claim: str, assertions: list[str], steps: list[str]) -> str:
    return next(iter(_claim_issues(claim, assertions, steps)), "")


def _check_assertions(check: dict[str, Any]) -> list[str] | None:
    assertions = check.get("assertions")
    if assertions is None and "assertion" in check:
        assertions = [check["assertion"]]
    if not isinstance(assertions, list) or not assertions or not all(
        isinstance(assertion, str) and assertion.startswith("ASSERT_") for assertion in assertions
    ):
        return None
    return [_normalize_generated_step(assertion) for assertion in assertions]


def _normalize_generated_step(step: str) -> str:
    parts = [part.strip() for part in step.split("|")]
    if parts[0].upper() in {
        "ASSERT_TEXT", "ASSERT_VALUE", "ASSERT_VALIDATION", "ASSERT_EACH_TEXT_CONTAINS",
        "ASSERT_DIALOG",
    } and len(parts) >= 2:
        value = parts[-1]
        if len(value) >= 2 and value[0] == value[-1] and value[0] in {"'", '"'}:
            parts[-1] = value[1:-1]
            return " | ".join(parts)
    return step


def _literal_forms(literal: str) -> set[str]:
    templated = re.sub(
        r"<([^>]+)>",
        lambda match: "${DATA:" + match.group(1).strip().casefold().replace(" ", "_") + "}",
        literal,
    )
    return {literal.casefold(), templated.casefold()}


def _source_prose_and_commands(source_step: str) -> tuple[str, list[str]]:
    lines = split_steps(source_step)
    commands = [line for line in lines if GenericStepExecutor.is_generic_step(line)]
    prose = "\n".join(line for line in lines if not GenericStepExecutor.is_generic_step(line))
    return prose, commands


def _temporal_issue(source_step: str, claims: list[str], checks: object, steps: list[str]) -> str:
    if any(step.startswith("ASSERT_DIALOG | ") for step in steps) and not any(
        step.startswith("WATCH_DIALOG | ") for step in steps
    ):
        return "ASSERT_DIALOG needs WATCH_DIALOG before the triggering action"
    prose, _ = _source_prose_and_commands(source_step)
    if re.search(r"\b(?:alert|dialog)\b", prose, re.I):
        watches = [i for i, step in enumerate(steps) if step.startswith("WATCH_DIALOG | ")]
        dialogs = [i for i, step in enumerate(steps) if step.startswith("ASSERT_DIALOG | ")]
        clicked = re.search(r"\bclick\s+'([^']+)'[^\n]*\b(?:alert|dialog)\b", prose, re.I)
        actions = [i for i, step in enumerate(steps) if step.startswith(("CLICK | ", "PRESS | "))
                   and (not clicked or clicked.group(1).casefold() in step.casefold())]
        if not any(watch < action < dialog for watch in watches for action in actions for dialog in dialogs):
            return "Browser alert needs WATCH_DIALOG before its action and ASSERT_DIALOG afterward"
    if not isinstance(checks, list):
        return ""
    for index, claim in enumerate(claims[1:], 1):
        after = re.search(r"\bafter\s+(?:clicking\s+)?([A-Za-z][\w ]*?)(?:,|\s+header\b|\s+the\b)", claim, re.I)
        if not after:
            continue
        action_name = after.group(1).strip().casefold()
        action_positions = [position for position, step in enumerate(steps)
                            if step.startswith("CLICK | ") and action_name in step.casefold()]
        earlier = _check_assertions(checks[index - 1]) if (
            index - 1 < len(checks) and isinstance(checks[index - 1], dict)
        ) else None
        if action_positions and earlier and any(steps.index(assertion) > action_positions[0] for assertion in earlier):
            return f"Assert the earlier outcome before clicking {after.group(1).strip()}"
    return ""


def _action_issue(source_step: str, steps: list[str]) -> str:
    prose, commands = _source_prose_and_commands(source_step)
    source = prose.casefold()
    clicks = [step.casefold() for step in steps if step.startswith("CLICK | ")]
    explicit_clicks = len(re.findall(r"\bclick\b", source)) + sum(
        step.startswith("CLICK | ") for step in commands
    )
    if len(clicks) < explicit_clicks:
        return f"Source has {explicit_clicks} click action(s), but generated steps have {len(clicks)}"
    if re.search(r"\badd\b[^\n]*\b(?:product|cart)\b", source) and not any(
        "add" in step for step in clicks
    ):
        return "Source add-to-cart action is missing from generated steps"
    if re.search(r"\b(?:remove|delete)\b[^\n]*\b(?:product|cart)\b|\bclick\s+'x'", source) and not any(
        any(word in step for word in ("remove", "delete", "cart_quantity_delete", "| x"))
        for step in clicks
    ):
        return "Source remove-product action is missing from generated steps"
    uploads = [step for step in steps if step.startswith("UPLOAD | ")]
    if ("upload" in source and not uploads) or len(uploads) < sum(
        step.startswith("UPLOAD | ") for step in commands
    ):
        return "Source file upload is missing from generated steps"
    fills = [step for step in steps if step.startswith("FILL | ")]
    if (re.search(r"\b(?:enter|fill)\b", source) and not fills) or len(fills) < sum(
        step.startswith("FILL | ") for step in commands
    ):
        return "Source form entry is missing from generated steps"
    return ""


def _scope_form_fields(source_step: str, steps: list[str]) -> list[str]:
    contexts: dict[str, str] = {}
    for line in split_steps(source_step):
        heading = re.search(r"\bunder\s+'([^']+)'", line, re.I)
        if not heading:
            continue
        before = line[:heading.start()].casefold()
        for field in ("name", "email", "password"):
            if re.search(rf"\b{field}\b", before):
                contexts[field] = heading.group(1)
    scoped = []
    for step in steps:
        parts = [part.strip() for part in step.split("|")]
        if parts[0].upper() == "FILL" and len(parts) == 3:
            field = parts[1].casefold().removesuffix(" address")
            if field in contexts:
                step = f"FILL | {parts[1]} in {contexts[field]} | {parts[2]}"
        scoped.append(step)
    headings = set(contexts.values())
    generic_submits = [index for index, step in enumerate(scoped)
                       if step.casefold() == "click | button[type='submit']"]
    click_labels = re.findall(r"\bclick\s+'([^']+)'", source_step, re.I)
    if len(headings) == len(generic_submits) == 1 and click_labels:
        heading = next(iter(headings))
        label = click_labels[-1]
        if label.casefold() in heading.casefold():
            scoped[generic_submits[0]] = f"CLICK | {label} in {heading}"
    return scoped


def _conversion_diagnostics(
    case: dict[str, Any], steps: list[str], checks: object,
) -> list[str]:
    """Collect independent defects so a focused retry sees more than the first failure."""
    issues: list[str] = []
    claims = _claims(case["expected"])
    mapped = checks if isinstance(checks, list) else []
    if len(mapped) != len(claims):
        issues.append("Every Expected Result clause needs an explicit assertion mapping")
    data = parse_test_data(case.get("data"))
    for index, claim in enumerate(claims):
        prefix = f"Clause {index + 1}: " if len(claims) > 1 else ""
        check = mapped[index] if index < len(mapped) else None
        if not isinstance(check, dict) or check.get("clause_index", index) != index:
            issues.append(prefix + "Expected Result clause has no executable assertion")
            continue
        assertions = _check_assertions(check)
        if not assertions or any(assertion not in steps for assertion in assertions):
            issues.append(prefix + "Expected Result clause has no executable assertion")
            continue
        issues.extend(prefix + issue for issue in _claim_issues(claim, assertions, steps))
        literals = re.findall(r"'([^']+)'", claim)
        if any(not any(
            any(form in assertion.casefold() for form in _literal_forms(literal)) or any(
                data.get(key.strip().casefold(), "").casefold() == literal.casefold()
                for key in re.findall(r"\$\{DATA:([^}]+)}", assertion, re.I)
            ) for assertion in assertions
        ) for literal in literals):
            issues.append(prefix + "A quoted expected value is missing from its assertion")
    action = _action_issue(case["step"], steps)
    if action:
        issues.append(action)
    ordering = _temporal_issue(case["step"], claims, checks, steps)
    if ordering:
        issues.append(ordering)
    unknown, placeholders = _data_key_issues(steps, case.get("data"))
    if unknown:
        issues.append("Generated command references an unknown Test Data field: " + ", ".join(unknown))
    if placeholders:
        issues.append("Generated command requires a real Test Data value: " + ", ".join(placeholders))
    for step in steps:
        parts = [part.strip() for part in step.split("|")]
        count = 3 if parts[0].upper() in THREE_ARGUMENTS else 2 if parts[0].upper() in TWO_ARGUMENTS else 1
        if len(parts) != count + 1 or not all(parts[1:]):
            issues.append(f"Invalid command arguments: {step}")
        elif parts[0].upper() == "UPLOAD" and len(parts) == 3:
            matches = re.findall(r"\$\{DATA:([^}]+)}", parts[2], re.I)
            value = data.get(matches[0].strip().casefold(), "") if matches else parts[2]
            if not Path(value).is_absolute() or not Path(value).is_file():
                issues.append("Generated upload requires an existing absolute file path")
    return list(dict.fromkeys(issues))


def validate_conversion(
    source: list[dict[str, Any]], payload: dict[str, Any], base_url: str,
    *, allow_site_actions: bool = False,
) -> list[dict[str, Any]]:
    """Keep every case, but quarantine incomplete or unsupported conversions."""
    from selfheal.onboarding import _suite_targets

    proposals = payload.get("cases")
    if not isinstance(proposals, list):
        raise ValueError("OpenAI conversion did not return a cases list")
    by_id: dict[str, dict[str, Any]] = {}
    expected_ids = {case["id"] for case in source}
    for proposal in proposals:
        if not isinstance(proposal, dict) or proposal.get("id") not in expected_ids:
            raise ValueError("OpenAI conversion returned an unknown test case")
        case_id = proposal["id"]
        if case_id in by_id:
            raise ValueError(f"OpenAI conversion duplicated {case_id}")
        by_id[case_id] = proposal
    if set(by_id) != expected_ids:
        raise ValueError("OpenAI conversion omitted test cases: " + ", ".join(sorted(expected_ids - set(by_id))))

    converted = []
    for case in source:
        proposal = by_id[case["id"]]
        raw_steps = proposal.get("steps")
        steps = _scope_form_fields(case["step"], [_normalize_generated_step(step) for step in raw_steps]) if (
            isinstance(raw_steps, list) and all(isinstance(s, str) for s in raw_steps)
        ) else []
        executable = {**case, "step": "\n".join(steps), "expected": ""}
        test_data = parse_test_data(case.get("data"))
        reason = str(proposal.get("review_reason") or "").strip()
        if not steps:
            reason = reason or "No executable commands were generated"
        claims = _claims(case["expected"])
        checks = proposal.get("checks")
        invalid_assertion = False
        if claims:
            if not isinstance(checks, list) or len(checks) != len(claims):
                reason = reason or "Every Expected Result clause needs an explicit assertion mapping"
                invalid_assertion = True
            else:
                for index, (claim, check) in enumerate(zip(claims, checks, strict=True)):
                    if not isinstance(check, dict) or (
                        check.get("clause_index", index) != index
                        or ("clause_index" not in check and check.get("expected") != claim)
                    ):
                        reason = reason or "Expected Result coverage does not match the source text"
                        invalid_assertion = True
                        break
                    assertions = _check_assertions(check)
                    if assertions is None or any(assertion not in steps for assertion in assertions):
                        reason = reason or "Expected Result clause has no executable assertion"
                        invalid_assertion = True
                        break
                    issue = _claim_issue(claim, assertions, steps)
                    if issue:
                        reason = reason or "Mapped assertion does not verify the stated result: " + issue
                        invalid_assertion = not missing_time_limit(claim) or any(
                            assertion.startswith("ASSERT_LOAD_TIME | ") for assertion in assertions
                        )
                        break
                    literals = re.findall(r"'([^']+)'", claim)
                    if any(
                        not any(
                            any(form in assertion.casefold() for form in _literal_forms(literal)) or any(
                                test_data.get(key.strip().casefold(), "").casefold() == literal.casefold()
                                for key in re.findall(r"\$\{DATA:([^}]+)}", assertion, re.I)
                            )
                            for assertion in assertions
                        )
                        for literal in literals
                    ):
                        reason = reason or "A quoted expected value is missing from its assertion"
                        invalid_assertion = True
                        break
                    for assertion in assertions:
                        assertion_parts = [part.strip() for part in assertion.split("|")]
                        if assertion_parts[0] in {
                            "ASSERT_TEXT", "ASSERT_VALUE", "ASSERT_VALIDATION",
                            "ASSERT_EACH_TEXT_CONTAINS", "ASSERT_DIALOG",
                        } and literals:
                            asserted_value = assertion_parts[-1].strip("'\"")
                            referenced = re.fullmatch(r"\$\{DATA:([^}]+)}", asserted_value, re.I)
                            if asserted_value.casefold() not in {
                                form for literal in literals for form in _literal_forms(literal)
                            } and not (
                                referenced and test_data.get(referenced.group(1).strip().casefold()) in literals
                            ) and not (
                                assertion_parts[0] == "ASSERT_VALIDATION" and asserted_value == "any"
                            ):
                                reason = reason or "Assert the quoted UI text, not the Expected Result description"
                                invalid_assertion = True
                                break
                    if invalid_assertion:
                        break
        ordering = _temporal_issue(case["step"], claims, checks, steps)
        if ordering:
            reason = reason or ordering
            invalid_assertion = True
        action = _action_issue(case["step"], steps)
        if action:
            reason = reason or action
            invalid_assertion = True
        syntax_error = ""
        if steps:
            try:
                _suite_targets([executable], base_url)
            except ValueError as exc:
                syntax_error = str(exc)
                reason = reason or syntax_error
        unknown_keys, placeholder_keys = _data_key_issues(steps, case.get("data"))
        unsafe = bool(SENSITIVE_ACTION.search(f"{case['step']} {case['expected']}")) and not allow_site_actions
        if unsafe:
            reason = "Sensitive site actions require explicit authorization"
        for step in steps:
            parts = [part.strip() for part in step.split("|")]
            count = 3 if parts[0].upper() in THREE_ARGUMENTS else 2 if parts[0].upper() in TWO_ARGUMENTS else 1
            if len(parts) != count + 1 or not all(parts[1:]):
                syntax_error = f"Invalid command arguments: {step}"
                reason = reason or syntax_error
                break
            if len(parts) > 1 and (
                re.search(r"\bdata-form-(?:action|id|class|heading)\b", parts[1], re.I)
                or ":contains(" in parts[1]
            ):
                syntax_error = "Generated selector uses snapshot-only metadata or unsupported CSS"
                reason = reason or syntax_error
                break
            if parts[0].upper() == "CLICK" and parts[1].startswith("/") and not parts[1].startswith("//"):
                syntax_error = "CLICK needs an element locator, not a URL path"
                reason = reason or syntax_error
                break
        for step in steps:
            parts = [part.strip() for part in step.split("|")]
            if "${ENV:" in step.upper():
                reason = "Generated commands cannot read environment secrets"
                unsafe = True
            elif parts[0].upper() == "UPLOAD" and not allow_site_actions:
                reason = "Generated file uploads require review"
                unsafe = True
            references = re.findall(r"\$\{DATA:([^}]+)}", step, re.I)
            if unknown_keys:
                reason = "Generated command references an unknown Test Data field: " + ", ".join(unknown_keys)
                unsafe = True
            elif placeholder_keys:
                reason = "Generated command requires a real Test Data value: " + ", ".join(placeholder_keys)
                unsafe = True
            elif not allow_site_actions and any(SENSITIVE_DATA_KEY.search(key) for key in references):
                reason = "Sensitive Test Data requires manual review before browser submission"
                unsafe = True
            elif parts[0].upper() == "UPLOAD" and len(parts) >= 3:
                value = (
                    test_data[references[0].strip().casefold()]
                    if references else parts[2]
                )
                upload = Path(value)
                if not upload.is_absolute() or not upload.is_file():
                    reason = "Generated upload requires an existing absolute file path"
                    unsafe = True
            elif parts[0].upper() in COLLECTION_ASSERTIONS:
                targets = parts[1:4] if parts[0].upper() == "ASSERT_MULTIPLY" else parts[1:3] if parts[0].upper() == "ASSERT_COUNT_MATCHES" else parts[1:2]
                if any(not _direct_reference(target) for target in targets):
                    reason = reason or "Collection assertions need explicit CSS or XPath locators"
                    syntax_error = reason
                if parts[0].upper() == "ASSERT_DATA_FIELDS":
                    keys = [key.strip().casefold() for key in parts[2].split(",")]
                    if len(keys) < 2 or any(key not in test_data or not test_data[key] for key in keys):
                        reason = "Address comparison needs at least two populated Test Data fields"
                        unsafe = True
            elif parts[0].upper() in {"FILL", "SELECT"} and references:
                target = parts[1].casefold()
                field_type = re.search(r"\b(email|password|address|name)\b", target)
                key = references[0].strip().casefold()
                if field_type and re.search(r"\b(email|password|address|name)\b", key):
                    if field_type.group(1) not in key:
                        reason = "Generated command maps a field to the wrong Test Data value"
                        unsafe = True
        diagnostics = _conversion_diagnostics(case, steps, checks)
        notes = [reason] if reason else []
        notes.extend(issue for issue in diagnostics if not any(issue in note for note in notes))
        reason = "; ".join(notes)
        has_assertion = any(step.strip().upper().startswith("ASSERT_") for step in steps)
        executable["execute_partial"] = (
            bool(reason) and has_assertion and not syntax_error and not unsafe and not invalid_assertion
        )
        executable["conversion_note"] = reason
        executable["conversion_issues"] = notes
        executable["original_step"] = case["step"]
        executable["original_expected"] = case["expected"]
        converted.append(executable)
    return converted


def _direct_reference(reference: str) -> bool:
    return reference.startswith(("css=", "xpath=", "#", ".", "[", "//"))


def save_converted_suite(
    source_path: Path, destination: Path, cases: list[dict[str, Any]],
    *, allow_site_actions: bool = False,
) -> None:
    """Persist executable rows and original source evidence together."""
    book = Workbook()
    sheet = book.active
    sheet.title = "Framework Cases"
    sheet.append([
        "TestCase ID", "Title", "Test Step", "Test Data", "Expected Result",
        "Conversion Note", "Original Test Step", "Original Expected Result",
    ])
    for case in cases:
        sheet.append([
            case["id"], case["title"], case["step"], case.get("data"), "",
            case["conversion_note"],
            case["original_step"], case["original_expected"],
        ])
    original = load_workbook(source_path, read_only=True, data_only=True)
    try:
        source_sheet = book.create_sheet("Original Cases")
        source = original["Original Cases"] if "Original Cases" in original.sheetnames else original.active
        for row in source.iter_rows(values_only=True):
            source_sheet.append(row)
    finally:
        original.close()
    meta = book.create_sheet("Framework Meta")
    meta.append(["Converter Version", CONVERTER_VERSION])
    meta.append(["Site Actions Authorized", allow_site_actions])
    meta.append(["Case ID", "Conversion Scope"])
    for case in cases:
        if case["conversion_note"] and case["execute_partial"]:
            meta.append([case["id"], "partial"])
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = destination.with_suffix(".xlsx.tmp")
    book.save(temporary)
    temporary.replace(destination)


def remove_generated_execute_column(path: Path) -> bool:
    """Migrate cached generated sheets without another conversion API call."""
    book = load_workbook(path)
    try:
        if "Original Cases" not in book.sheetnames or "Framework Meta" not in book.sheetnames:
            return False
        sheet = book["Framework Cases"] if "Framework Cases" in book.sheetnames else book.active
        headers = [cell.value for cell in sheet[1]]
        if "Execute" not in headers:
            return False
        execute_column = headers.index("Execute") + 1
        note_column = headers.index("Conversion Note") + 1
        id_column = headers.index("TestCase ID") + 1
        meta = book["Framework Meta"]
        existing_partial = {
            str(row[0]).casefold()
            for row in meta.iter_rows(min_row=4, values_only=True)
            if len(row) >= 2 and row[0] and str(row[1] or "").casefold() == "partial"
        }
        if meta.max_row < 3:
            meta.append(["Case ID", "Conversion Scope"])
        for row in sheet.iter_rows(min_row=2):
            case_id = str(row[id_column - 1].value or "")
            execute = str(row[execute_column - 1].value or "").strip().casefold()
            note = str(row[note_column - 1].value or "").strip()
            if case_id and note and execute in {"yes", "true"} and case_id.casefold() not in existing_partial:
                meta.append([case_id, "partial"])
                existing_partial.add(case_id.casefold())
        sheet.delete_cols(execute_column)
        temporary = path.with_suffix(".xlsx.tmp")
        book.save(temporary)
        temporary.replace(path)
        return True
    finally:
        book.close()


def review_cases(path: Path) -> list[dict[str, Any]]:
    from selfheal.suite_loader import read_excel_tests

    enabled_rows = {case["row"] for case in read_excel_tests(path)}
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = book.active.iter_rows(values_only=True)
        headers = next(rows)
        reviews = []
        for number, row in enumerate(rows, 2):
            values = dict(zip(headers, row, strict=False))
            note = str(values.get("Conversion Note") or "").strip()
            if not note or number in enabled_rows:
                continue
            if note in {
                "Generated command references an unknown Test Data field",
                "Generated command requires a real Test Data value",
            }:
                unknown, placeholders = _data_key_issues(
                    split_steps(values.get("Test Step")), values.get("Test Data"),
                )
                keys = unknown if "unknown" in note else placeholders
                if keys:
                    note += ": " + ", ".join(keys)
            reviews.append({
                "id": str(values["TestCase ID"]),
                "title": str(values.get("Title") or ""),
                "row": number,
                "reason": note,
            })
        return reviews
    finally:
        book.close()


def reviewed_case_ids(path: Path) -> list[str]:
    """Include blocked and partially executable generated rows."""
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        if "Original Cases" not in book.sheetnames:
            return []
        sheet = book["Framework Cases"]
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows)
        return [
            str(values["TestCase ID"])
            for row in rows
            if (values := dict(zip(headers, row, strict=False))).get("TestCase ID")
            and str(values.get("Conversion Note") or "").strip()
        ]
    finally:
        book.close()


def replace_reviewed_conversions(path: Path, cases: list[dict[str, Any]]) -> int:
    """Replace only reviewed rows; preserve accepted commands and their data."""
    replacements = {case["id"]: case for case in cases}
    if not replacements:
        return 0
    book = load_workbook(path)
    temporary = path.with_name(f".{path.stem}.{uuid4().hex}.xlsx")
    try:
        sheet = book["Framework Cases"]
        headers = [cell.value for cell in sheet[1]]
        columns = {name: headers.index(name) + 1 for name in (
            "TestCase ID", "Test Step", "Expected Result", "Conversion Note",
        )}
        replaced: set[str] = set()
        for row in range(2, sheet.max_row + 1):
            case_id = str(sheet.cell(row, columns["TestCase ID"]).value or "")
            if case_id not in replacements or not str(sheet.cell(row, columns["Conversion Note"]).value or "").strip():
                continue
            case = replacements[case_id]
            sheet.cell(row, columns["Test Step"]).value = case["step"]
            sheet.cell(row, columns["Expected Result"]).value = ""
            sheet.cell(row, columns["Conversion Note"]).value = case["conversion_note"]
            replaced.add(case_id)
        if not replaced:
            return 0
        meta = book["Framework Meta"]
        for row in range(meta.max_row, 3, -1):
            if str(meta.cell(row, 1).value or "") in replaced:
                meta.delete_rows(row)
        for case_id in replaced:
            case = replacements[case_id]
            if case["conversion_note"] and case["execute_partial"]:
                meta.append([case_id, "partial"])
        book.save(temporary)
        temporary.replace(path)
        return len(replaced)
    finally:
        book.close()
        temporary.unlink(missing_ok=True)


def reuse_valid_conversions(
    path: Path, *, source: Path, base_url: str, allow_site_actions: bool = False,
) -> int:
    """Revalidate accepted single-claim rows from an older converter cache."""
    from selfheal.suite_loader import read_excel_tests

    current_sources = {case["id"]: case for case in read_excel_tests(source)}
    current_pending = set(reviewed_case_ids(path))
    if not current_pending:
        return 0
    restored: dict[str, dict[str, Any]] = {}
    for old_path in sorted(path.parent.glob(f"{source.stem}-*.xlsx"), reverse=True):
        if old_path == path or not is_converted_suite(old_path):
            continue
        old_source = {case["id"]: case for case in read_excel_tests(old_path, "Original Cases")}
        for old_case in read_excel_tests(old_path):
            case_id = old_case["id"]
            original = old_source.get(case_id)
            current = current_sources.get(case_id)
            if (
                case_id not in current_pending or case_id in restored
                or old_case.get("conversion_note") or not original or not current
                or any(original.get(key) != current.get(key) for key in ("step", "data", "expected"))
                or len(_claims(current["expected"])) != 1
            ):
                continue
            steps = split_steps(old_case["step"])
            assertions = [step for step in steps if step.startswith("ASSERT_")]
            if not assertions or any(PRIVATE_VALUE.search(step) or PRIVATE_ASSIGNMENT.search(step) for step in steps):
                continue
            proposal = {"cases": [{"id": case_id, "steps": steps,
                                   "checks": [{"clause_index": 0, "assertions": assertions}]}]}
            validated = validate_conversion(
                [current], proposal, base_url, allow_site_actions=allow_site_actions,
            )[0]
            if not validated["conversion_note"]:
                restored[case_id] = validated
    return replace_reviewed_conversions(path, list(restored.values()))


def recover_passed_conversions(
    path: Path, *, root: Path, base_url: str, allow_site_actions: bool = False,
) -> int:
    """Restore validated commands for same-source cases passed in an earlier run."""
    from selfheal.suite_loader import read_excel_tests

    pending = {item["id"] for item in review_cases(path)}
    if not pending:
        return 0
    sources = {
        case["id"]: case for case in read_excel_tests(path, "Original Cases")
        if case["id"] in pending and len(_claims(case["expected"])) == 1
    }
    if not sources:
        return 0

    recovered: dict[str, str] = {}
    for result_file in sorted((root / "data" / "runs").glob("run_*/results.json"), reverse=True):
        if len(recovered) == len(sources):
            break
        try:
            result = json.loads(result_file.read_text(encoding="utf-8"))
            previous_suite = Path(result.get("suite", ""))
            if not previous_suite.is_absolute():
                previous_suite = root / previous_suite
            if previous_suite.resolve() != path.resolve() or result.get("base_url") != base_url:
                continue
        except (OSError, ValueError, TypeError):
            continue
        for test in result.get("tests", []):
            case_id = test.get("id")
            if case_id not in sources or case_id in recovered or test.get("status") != "passed":
                continue
            executed = test.get("steps", [])
            if not executed or any(step.get("status") != "passed" for step in executed):
                continue
            steps = [step.get("text") for step in executed]
            if not all(isinstance(step, str) for step in steps):
                continue
            if any(PRIVATE_VALUE.search(step) or PRIVATE_ASSIGNMENT.search(step) for step in steps):
                continue
            assertions = [step for step in steps if step.strip().upper().startswith("ASSERT_")]
            if len(assertions) != 1:
                continue
            proposal = {"cases": [{
                "id": case_id, "steps": steps,
                "checks": [{"clause_index": 0, "assertion": assertions[0]}],
            }]}
            try:
                checked = validate_conversion(
                    [sources[case_id]], proposal, base_url,
                    allow_site_actions=allow_site_actions,
                )[0]
            except ValueError:
                continue
            if not checked["conversion_note"]:
                recovered[case_id] = checked["step"]

    if not recovered:
        return 0
    book = load_workbook(path)
    temporary = path.with_name(f".{path.stem}.{uuid4().hex}.xlsx")
    try:
        sheet = book["Framework Cases"]
        headers = [cell.value for cell in sheet[1]]
        columns = {name: headers.index(name) + 1 for name in (
            "TestCase ID", "Test Step", "Expected Result", "Conversion Note",
        )}
        for row in range(2, sheet.max_row + 1):
            case_id = sheet.cell(row, columns["TestCase ID"]).value
            if case_id not in recovered:
                continue
            sheet.cell(row, columns["Test Step"]).value = recovered[case_id]
            sheet.cell(row, columns["Expected Result"]).value = ""
            sheet.cell(row, columns["Conversion Note"]).value = ""
        book.save(temporary)
        temporary.replace(path)
    finally:
        book.close()
        temporary.unlink(missing_ok=True)
    return len(recovered)
