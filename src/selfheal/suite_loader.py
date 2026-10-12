"""Read spreadsheet test cases into the runner's normalized format."""

from __future__ import annotations

from pathlib import Path
from typing import Any

REQUIRED_COLUMNS = {"TestCase ID", "Title", "Test Step", "Test Data", "Expected Result"}


def read_excel_tests(file_path: str | Path, sheet_name: str | None = None) -> list[dict[str, Any]]:
    """Return enabled test cases from a workbook, preserving spreadsheet row numbers."""
    import openpyxl

    workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name] if sheet_name else workbook.active
        generated = sheet.title == "Framework Cases" and "Original Cases" in workbook.sheetnames
        partial_ids: set[str] = set()
        if generated and "Framework Meta" in workbook.sheetnames:
            partial_ids = {
                str(row[0]).casefold()
                for row in workbook["Framework Meta"].iter_rows(min_row=4, values_only=True)
                if len(row) >= 2 and row[0] and str(row[1] or "").casefold() == "partial"
            }
        header_row = next(sheet.iter_rows(min_row=1, max_row=1))
        headers = [str(cell.value or "").strip() for cell in header_row]
        missing = REQUIRED_COLUMNS.difference(headers)
        if missing:
            raise ValueError("Missing Excel columns: " + ", ".join(sorted(missing)))

        tests: list[dict[str, Any]] = []
        for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            values = dict(zip(headers, row, strict=False))
            if not values.get("TestCase ID"):
                continue
            legacy_execute = str(values.get("Execute") or "").strip().casefold()
            if legacy_execute in {"no", "skip", "false"}:
                continue
            note = str(values.get("Conversion Note") or "").strip()
            if generated and note and legacy_execute not in {"yes", "true"} and (
                str(values["TestCase ID"]).casefold() not in partial_ids
            ):
                continue
            tests.append(
                {
                    "id": str(values["TestCase ID"]),
                    "title": str(values.get("Title") or ""),
                    "step": str(values.get("Test Step") or ""),
                    "data": values.get("Test Data"),
                    "expected": str(values.get("Expected Result") or ""),
                    "conversion_note": note,
                    "row": row_number,
                }
            )
        return tests
    finally:
        workbook.close()
